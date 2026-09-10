from __future__ import annotations

import asyncio
import base64
import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import stat
import subprocess
import sys
import threading
import time
import tracemalloc
from collections.abc import Iterator
from contextlib import contextmanager
from datetime import timedelta
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pypdfium2 as pdfium
import pytest
import stripe
from bs4 import BeautifulSoup
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from unrender.product import admin
from unrender.product import backup as backup_module
from unrender.product import service as service_module
from unrender.product.backup import BackupError, create_backup, restore_backup
from unrender.product.config import Settings
from unrender.product.database import SCHEMA_VERSION, Database
from unrender.product.extractors import (
    ExtractionError,
    ExtractionOutput,
    ModalExtractor,
    ReplayExtractor,
)
from unrender.product.security import hash_password, password_needs_rehash, verify_password
from unrender.product.service import (
    ProductError,
    ProductService,
    StorageReservation,
    timestamp,
    utcnow,
)
from unrender.product.storage import InvalidUpload, Storage
from unrender.product.web import (
    CSRF_COOKIE,
    BodyLimitMiddleware,
    SecurityHeadersMiddleware,
    create_app,
)
from unrender.product.worker import JobWorker
from unrender.schema.chart_schema import ChartData

STATIC_DIR = Path(__file__).parents[1] / "unrender" / "product" / "static"


def settings_for(tmp_path: Path, **changes: object) -> Settings:
    values: dict[str, object] = {
        "data_dir": tmp_path / "data",
        "environment": "test",
        "base_url": "http://testserver",
        "extractor_backend": "replay",
        "worker_enabled": False,
        "allow_registration": True,
        "seed_demo_account": True,
        "initial_credits": 3,
        "rate_limit_per_minute": 1000,
    }
    values.update(changes)
    return Settings(**values)  # type: ignore[arg-type]


def service_for(tmp_path: Path, **changes: object) -> ProductService:
    settings = settings_for(tmp_path, **changes)
    service = ProductService(
        settings=settings,
        database=Database(settings.database_path),
        storage=Storage(settings),
        extractor=ReplayExtractor(STATIC_DIR),
        static_dir=STATIC_DIR,
    )
    service.initialize()
    return service


def customer_id(service: ProductService, email: str = "customer@example.com") -> str:
    session = service.register(email, "customer password is long enough")
    user = service.session_user(session["session"])
    assert user is not None
    return str(user["id"])


def csrf_headers(client: TestClient) -> dict[str, str]:
    return {"X-CSRF-Token": client.cookies.get(CSRF_COOKIE)}


def register_and_create_api_key(
    client: TestClient,
    *,
    email: str = "api@example.com",
) -> str:
    registered = client.post(
        "/api/auth/register",
        json={"email": email, "password": "api customer password is long enough"},
    )
    assert registered.status_code == 201
    created = client.post(
        "/api/keys",
        headers=csrf_headers(client),
        json={"name": "integration"},
    )
    assert created.status_code == 201
    return str(created.json()["key"])


def png_bytes(*, color: str = "white") -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (160, 100), color=color).save(output, format="PNG")
    return output.getvalue()


def test_configuration_rejects_unsafe_production_and_live_billing(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="HTTPS"):
        settings_for(
            tmp_path,
            environment="production",
            base_url="http://example.com",
            extractor_backend="modal",
            seed_demo_account=False,
        ).validate()
    with pytest.raises(ValueError, match="Replay"):
        settings_for(
            tmp_path,
            environment="production",
            base_url="https://example.com",
            seed_demo_account=False,
        ).validate()
    with pytest.raises(ValueError, match="MODAL_REVISION"):
        settings_for(
            tmp_path,
            environment="production",
            base_url="https://example.com",
            extractor_backend="modal",
            seed_demo_account=False,
            allow_registration=False,
            worker_enabled=True,
            modal_model_path="approved/unrender-model",
        ).validate()
    with pytest.raises(ValueError, match="ALLOW_REGISTRATION"):
        settings_for(
            tmp_path,
            environment="production",
            base_url="https://example.com",
            extractor_backend="modal",
            seed_demo_account=False,
            modal_model_path="approved/unrender-model",
            modal_model_revision="0123456789abcdef0123456789abcdef01234567",
        ).validate()
    with pytest.raises(ValueError, match="MODEL must be"):
        settings_for(
            tmp_path,
            environment="production",
            base_url="https://example.com",
            extractor_backend="modal",
            seed_demo_account=False,
            allow_registration=False,
            worker_enabled=True,
            modal_model_revision="0123456789abcdef0123456789abcdef01234567",
        ).validate()
    settings_for(
        tmp_path,
        environment="production",
        base_url="https://example.com",
        extractor_backend="modal",
        seed_demo_account=False,
        allow_registration=False,
        modal_model_path="approved/unrender-model",
        modal_model_revision="0123456789abcdef0123456789abcdef01234567",
        modal_model_digest="b" * 64,
        modal_provider_release="a" * 64,
        worker_enabled=True,
    ).validate()
    with pytest.raises(ValueError, match="PROVIDER_RELEASE"):
        settings_for(
            tmp_path,
            environment="production",
            base_url="https://example.com",
            extractor_backend="modal",
            seed_demo_account=False,
            allow_registration=False,
            worker_enabled=True,
            modal_model_path="approved/unrender-model",
            modal_model_revision="0123456789abcdef0123456789abcdef01234567",
            modal_model_digest="b" * 64,
        ).validate()
    with pytest.raises(ValueError, match="MODEL_DIGEST"):
        settings_for(
            tmp_path,
            environment="production",
            base_url="https://example.com",
            extractor_backend="modal",
            seed_demo_account=False,
            allow_registration=False,
            worker_enabled=True,
            modal_model_path="approved/unrender-model",
            modal_model_revision="0123456789abcdef0123456789abcdef01234567",
            modal_provider_release="a" * 64,
        ).validate()
    with pytest.raises(ValueError, match="WORKER_ENABLED"):
        settings_for(
            tmp_path,
            environment="production",
            base_url="https://example.com",
            extractor_backend="modal",
            seed_demo_account=False,
            allow_registration=False,
            worker_enabled=False,
            modal_model_path="approved/unrender-model",
            modal_model_revision="0123456789abcdef0123456789abcdef01234567",
        ).validate()
    with pytest.raises(ValueError, match="MODAL_FUNCTION"):
        settings_for(
            tmp_path,
            environment="production",
            base_url="https://example.com",
            extractor_backend="modal",
            seed_demo_account=False,
            allow_registration=False,
            worker_enabled=True,
            modal_function_name="infer-one",
            modal_model_path="approved/unrender-model",
            modal_model_revision="0123456789abcdef0123456789abcdef01234567",
            modal_model_digest="b" * 64,
            modal_provider_release="a" * 64,
        ).validate()
    with pytest.raises(ValueError, match="test-mode"):
        settings_for(
            tmp_path,
            stripe_secret_key="sk_live_forbidden",
            stripe_webhook_secret="whsec_test",
            stripe_price_id="price_test",
        ).validate()
    with pytest.raises(ValueError, match="secret, webhook secret, and price"):
        settings_for(tmp_path, stripe_secret_key="sk_test_partial").validate()
    with pytest.raises(ValueError, match="without a path"):
        settings_for(tmp_path, base_url="https://example.com/unrender").validate()
    with pytest.raises(ValueError, match="MAX_USER_STORAGE_BYTES"):
        settings_for(
            tmp_path,
            max_upload_bytes=1024,
            max_user_storage_bytes=512,
        ).validate()
    with pytest.raises(ValueError, match="Tenant upload limits"):
        settings_for(tmp_path, max_unattached_uploads=0).validate()


def test_runtime_lock_includes_the_production_provider_client() -> None:
    lock = (STATIC_DIR.parents[2] / "requirements-app.lock").read_text(encoding="utf-8")
    assert "\nmodal==" in lock


def test_modal_release_handshake_rejects_drift_and_records_approved_release(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fixture = json.loads(
        (STATIC_DIR / "demo" / "budget-quarter-result.json").read_text(encoding="utf-8")
    )
    release = "a" * 64
    response = {
        "json": fixture["result"],
        "raw": fixture["raw"],
        "provider_release": "b" * 64,
    }
    remote_calls: list[tuple[object, ...]] = []

    async def remote_call(*args: object) -> dict[str, object]:
        remote_calls.append(args)
        return response

    remote = SimpleNamespace(remote=SimpleNamespace(aio=remote_call))
    modal = SimpleNamespace(
        Function=SimpleNamespace(from_name=lambda *_: remote),
    )
    monkeypatch.setitem(sys.modules, "modal", modal)
    extractor = ModalExtractor(
        settings_for(
            tmp_path,
            extractor_backend="modal",
            modal_model_path="approved/unrender-model",
            modal_model_revision="1" * 40,
            modal_model_digest="2" * 64,
            modal_provider_release=release,
        )
    )
    with pytest.raises(ExtractionError, match="approved deployment"):
        extractor.extract(png_bytes())

    response["provider_release"] = release
    output = extractor.extract(png_bytes())
    assert remote_calls[-1][1:] == ("approved/unrender-model", "1" * 40, "2" * 64)
    assert output.model_version.endswith(f"+provider:{release[:12]}")


def test_passwords_are_salted_and_verified() -> None:
    first = hash_password("correct horse battery staple")
    second = hash_password("correct horse battery staple")
    assert first != second
    assert verify_password("correct horse battery staple", first)
    assert not verify_password("wrong password", first)
    assert first.split("$")[3] == "5"
    assert not password_needs_rehash(first)


def test_login_opportunistically_upgrades_legacy_scrypt_parameters(tmp_path: Path) -> None:
    service = service_for(tmp_path, seed_demo_account=False)
    user_id = service.provision_user("legacy@example.com", "legacy password is long enough")
    salt = b"legacy-salt-1234"
    derived = hashlib.scrypt(b"legacy password is long enough", salt=salt, n=2**14, r=8, p=1)
    legacy = "scrypt$16384$8$1$" + base64.urlsafe_b64encode(salt).decode("ascii")
    legacy += "$" + base64.urlsafe_b64encode(derived).decode("ascii")
    with service.database.transaction(immediate=True) as conn:
        conn.execute("UPDATE users SET password_hash=? WHERE id=?", (legacy, user_id))

    service.authenticate("legacy@example.com", "legacy password is long enough")

    with service.database.connect() as conn:
        upgraded = conn.execute("SELECT password_hash FROM users WHERE id=?", (user_id,)).fetchone()
    assert upgraded["password_hash"].split("$")[3] == "5"
    assert not password_needs_rehash(upgraded["password_hash"])


def test_operator_can_provision_when_public_registration_is_closed(tmp_path: Path) -> None:
    service = service_for(tmp_path, allow_registration=False, seed_demo_account=False)
    with pytest.raises(ProductError, match="closed"):
        service.register("blocked@example.com", "a sufficiently long password")
    user_id = service.provision_user(
        "invited@example.com", "an operator supplied password", credits=7
    )
    assert service.account(user_id)["credits"] == 7
    with service.database.connect() as conn:
        ledger = conn.execute(
            "SELECT reason FROM credit_ledger WHERE user_id=?", (user_id,)
        ).fetchone()
    assert ledger["reason"] == "operator_grant"
    session = service.authenticate("invited@example.com", "an operator supplied password")
    assert service.session_user(session["session"])["id"] == user_id


def test_cancel_wins_atomic_race_with_worker_completion(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = service_for(tmp_path)
    user_id = customer_id(service)
    upload = service.prepare_upload(
        user_id=user_id, filename="customer-chart.png", content=png_bytes(color="blue")
    )
    job = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    fixture = json.loads(
        (STATIC_DIR / "demo" / "budget-quarter-result.json").read_text(encoding="utf-8")
    )
    service.extractor = SimpleNamespace(
        extract=lambda _: ExtractionOutput(
            chart=ChartData.model_validate(fixture["result"]),
            raw="atomic completion test",
            extractor="test",
            model_version="test-pinned",
        )
    )

    original_transaction = service.database.transaction
    completion_waiting = threading.Event()
    allow_completion = threading.Event()
    transaction_count = 0
    count_lock = threading.Lock()

    @contextmanager
    def gated_transaction(*, immediate: bool = False) -> Iterator[object]:
        nonlocal transaction_count
        with count_lock:
            transaction_count += 1
            current_count = transaction_count
        if current_count == 4:
            completion_waiting.set()
            assert allow_completion.wait(timeout=5)
        with original_transaction(immediate=immediate) as conn:
            yield conn

    monkeypatch.setattr(service.database, "transaction", gated_transaction)
    worker = threading.Thread(target=service.process_one)
    worker.start()
    assert completion_waiting.wait(timeout=5)
    service.cancel(user_id=user_id, job_id=job["id"])
    allow_completion.set()
    worker.join(timeout=5)

    assert not worker.is_alive()
    assert service.get_job(user_id=user_id, job_id=job["id"])["status"] == "cancelled"
    assert service.account(user_id)["credits"] == 3


def test_cancel_request_wins_when_provider_fails(tmp_path: Path) -> None:
    service = service_for(tmp_path)
    user_id = customer_id(service)
    upload = service.prepare_upload(
        user_id=user_id, filename="customer-chart.png", content=png_bytes(color="blue")
    )
    job = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    provider_started = threading.Event()
    provider_release = threading.Event()

    def fail_after_cancel(_: bytes) -> None:
        provider_started.set()
        assert provider_release.wait(timeout=5)
        raise ExtractionError("provider_unavailable", "provider failed")

    service.extractor = SimpleNamespace(extract=fail_after_cancel)
    worker = threading.Thread(target=service.process_one)
    worker.start()
    assert provider_started.wait(timeout=5)
    service.cancel(user_id=user_id, job_id=job["id"])
    provider_release.set()
    worker.join(timeout=5)

    assert not worker.is_alive()
    assert service.get_job(user_id=user_id, job_id=job["id"])["status"] == "cancelled"
    assert service.account(user_id)["credits"] == 2


def test_provider_failures_emit_privacy_safe_operational_fields(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    service = service_for(tmp_path, seed_demo_account=False)
    user_id = customer_id(service)
    upload = service.prepare_upload(
        user_id=user_id,
        filename="customer-private-name.png",
        content=png_bytes(color="purple"),
    )
    job = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)

    def fail_provider(_: bytes) -> None:
        raise ExtractionError("provider_unavailable", "secret upstream detail")

    service.extractor = SimpleNamespace(extract=fail_provider)
    with caplog.at_level("INFO", logger="unrender.product"):
        assert service.process_one()

    failed = next(record for record in caplog.records if record.message == "provider_call_failed")
    assert failed.event_name == "provider_call_failed"
    assert failed.error_code == "provider_unavailable"
    assert failed.job_id == job["id"]
    assert isinstance(failed.duration_ms, int)
    rendered = " ".join(record.getMessage() for record in caplog.records)
    assert "customer-private-name" not in rendered
    assert "secret upstream detail" not in rendered


def test_reprocess_failure_and_cancel_preserve_approved_result(tmp_path: Path) -> None:
    service = service_for(tmp_path)
    user_id = customer_id(service)
    upload = service.prepare_upload(
        user_id=user_id, filename="customer-chart.png", content=png_bytes(color="blue")
    )
    job = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    fixture = json.loads(
        (STATIC_DIR / "demo" / "budget-quarter-result.json").read_text(encoding="utf-8")
    )
    service.extractor = SimpleNamespace(
        extract=lambda _: ExtractionOutput(
            chart=ChartData.model_validate(fixture["result"]),
            raw="approved result preservation test",
            extractor="test",
            model_version="test-pinned",
        )
    )
    assert service.process_one()
    approved = service.approve(user_id=user_id, job_id=job["id"])
    approved_result = approved["result"]
    approved_at = approved["approved_at"]
    assert service.account(user_id)["credits"] == 2

    service.reprocess(user_id=user_id, job_id=job["id"])

    def fail_provider(_: bytes) -> None:
        raise ExtractionError("provider_unavailable", "provider failed")

    service.extractor = SimpleNamespace(extract=fail_provider)
    assert service.process_one()
    after_failure = service.get_job(user_id=user_id, job_id=job["id"])
    assert after_failure["status"] == "approved"
    assert after_failure["approved_at"] == approved_at
    assert after_failure["result"] == approved_result
    assert after_failure["error"]["code"] == "provider_unavailable"
    assert service.account(user_id)["credits"] == 1

    service.reprocess(user_id=user_id, job_id=job["id"])
    after_cancel = service.cancel(user_id=user_id, job_id=job["id"])
    assert after_cancel["status"] == "approved"
    assert after_cancel["approved_at"] == approved_at
    assert after_cancel["result"] == approved_result
    assert service.account(user_id)["credits"] == 1


def test_admin_cli_prompts_for_password_and_creates_account(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    data_dir = tmp_path / "operator-data"
    monkeypatch.setenv("UNRENDER_DATA_DIR", str(data_dir))
    monkeypatch.setenv("UNRENDER_ENV", "test")
    monkeypatch.setenv("UNRENDER_BASE_URL", "http://testserver")
    monkeypatch.setenv("UNRENDER_SEED_DEMO", "false")
    monkeypatch.setenv("UNRENDER_ALLOW_REGISTRATION", "false")
    monkeypatch.setattr(
        "sys.argv", ["unrender-admin", "create-user", "operator@example.com", "--credits", "4"]
    )
    passwords = iter(["an operator supplied password", "an operator supplied password"])
    monkeypatch.setattr(admin.getpass, "getpass", lambda _: next(passwords))

    admin.main()

    assert capsys.readouterr().out.startswith("Created account ")
    with Database(data_dir / "unrender.sqlite3").connect() as conn:
        row = conn.execute(
            "SELECT email, credit_balance FROM users WHERE email=?", ("operator@example.com",)
        ).fetchone()
    assert dict(row) == {"email": "operator@example.com", "credit_balance": 4}


def test_storage_validates_magic_pdf_limits_and_crop(tmp_path: Path) -> None:
    storage = Storage(settings_for(tmp_path, max_pdf_pages=1))
    with pytest.raises(InvalidUpload, match="PNG, JPEG, WebP, or PDF"):
        storage.inspect(b"not-an-image")

    document = pdfium.PdfDocument.new()
    document.new_page(160, 100)
    document.new_page(160, 100)
    pdf_buffer = io.BytesIO()
    document.save(pdf_buffer)
    pdf = pdf_buffer.getvalue()
    document.close()
    with pytest.raises(InvalidUpload, match="limit is 1"):
        storage.inspect(pdf)

    one_page = pdfium.PdfDocument.new()
    one_page.new_page(160, 100)
    one_page_buffer = io.BytesIO()
    one_page.save(one_page_buffer)
    one_page.close()
    pdf_content = one_page_buffer.getvalue()
    pdf_inspection = storage.inspect(pdf_content)
    pdf_source = storage.save_upload(
        user_id="pdf-user", upload_id="pdf-upload", content=pdf_content
    )
    rendered = storage.page_png(
        source=pdf_source,
        mime_type=pdf_inspection.mime_type,
        page_index=0,
    )
    assert rendered.startswith(b"\x89PNG")

    encrypted_pdf = base64.b64decode(
        "JVBERi0xLjcKJcK1wrYKJSBXcml0dGVuIGJ5IE11UERGIDEuMjguMgoKMSAwIG9iago8PC9UeXBlL0NhdGFsb2cvUGFnZXMgMiAwIFIvSW5mbzw8L1Byb2R1Y2VyPDcxMjRBN0RCNjg1MEZCQkVEMzI5NUUzM0VDRUYxOTQ1MkQxQ0U2OEI5QkM0QUZGRDJGNjhEN0NDQ0I5NERGREQ+Pj4+PgplbmRvYmoKCjIgMCBvYmoKPDwvVHlwZS9QYWdlcy9Db3VudCAxL0tpZHNbNCAwIFJdPj4KZW5kb2JqCgozIDAgb2JqCjw8Pj4KZW5kb2JqCgo0IDAgb2JqCjw8L1R5cGUvUGFnZS9NZWRpYUJveFswIDAgNzIgNzJdL1JvdGF0ZSAwL1Jlc291cmNlcyAzIDAgUi9QYXJlbnQgMiAwIFI+PgplbmRvYmoKCnhyZWYKMCA1CjAwMDAwMDAwMDAgNjU1MzUgZiAKMDAwMDAwMDA0MiAwMDAwMCBuIAowMDAwMDAwMTcyIDAwMDAwIG4gCjAwMDAwMDAyMjQgMDAwMDAgbiAKMDAwMDAwMDI0NSAwMDAwMCBuIAoKdHJhaWxlcgo8PC9TaXplIDUvUm9vdCAxIDAgUi9JRFs8QzNBQUMzODMxQkMyOTdDMjk1QzI5RDAyQzI5RDQ3QzI+PDIwMjNGOUJCMzYwNEUyMENFN0YwMUJDQjU3NDREM0ZEPl0vRW5jcnlwdDw8L0ZpbHRlci9TdGFuZGFyZC9SIDYvViA1L0xlbmd0aCAyNTYvUCAtNC9FbmNyeXB0TWV0YWRhdGEgdHJ1ZS9TdG1GL1N0ZENGL1N0ckYvU3RkQ0YvQ0Y8PC9TdGRDRjw8L0F1dGhFdmVudC9Eb2NPcGVuL0NGTS9BRVNWMy9MZW5ndGggMzI+Pj4+L088RTExQkE5N0Q5NEYyQkFFMzNGQ0IzQUM5M0QxQjBFQzdDODlCMEU5RTUwQUE3Q0M1NDMzRDRCODBCMkNFQzVGNUFBQTlCODU2MzkzRTFDRDQxMjY3MTk2Q0I0RDFDOEM1Pi9VPDkyRTVGMTZFOTM4RjJDMDA0OUYwRjIzODM2NDEyMzdFRjIwMDNGRDc4RTlDREQwQjFBRTIxMTUwREM3NjEyMTI0NkIxNUNGMEFGRDg0MDQ2QTc4RTAzRjBFN0ZCNTUzQz4vT0U8MUVGMUNGRjdCMzVBQ0Q2MTk2M0JCQTYyRDdBRDg5MkNGMTI5NDgxQkE5MUNBRkIwRDY4MkI4NTNDNzdENkZDRT4vVUU8OUM3NjY3NzJERkE3MzVDQTY5MUFFODczMEU2NENDQkVEQkI5MjlENTY0RDI3RDI0RjI5MjY0RDFBMTMwQjQ4MD4vUGVybXM8NDgxNEJDRjc0MDExOEFBQ0FGMzBDRUEzRDcwOEI3RjE+Pj4+PgpzdGFydHhyZWYKMzM0CiUlRU9GCg=="
    )
    with pytest.raises(InvalidUpload, match="Password-protected"):
        storage.inspect(encrypted_pdf)

    image = png_bytes()
    inspection = storage.inspect(image)
    source = storage.save_upload(user_id="user", upload_id="upload", content=image)
    cropped = storage.page_png(
        source=source,
        mime_type=inspection.mime_type,
        page_index=0,
        crop={"x": 0.1, "y": 0.1, "width": 0.5, "height": 0.5},
    )
    assert cropped.startswith(b"\x89PNG")
    with pytest.raises(InvalidUpload, match="within"):
        storage.page_png(
            source=source,
            mime_type=inspection.mime_type,
            page_index=0,
            crop={"x": 0.8, "y": 0.1, "width": 0.5, "height": 0.5},
        )


def test_database_and_customer_sources_are_private_on_disk(tmp_path: Path) -> None:
    service = service_for(tmp_path, seed_demo_account=False)
    user_id = customer_id(service)
    upload = service.prepare_upload(
        user_id=user_id,
        filename="private.png",
        content=png_bytes(),
    )
    row = service.upload(user_id=user_id, upload_id=upload["id"])

    assert stat.S_IMODE(service.settings.data_dir.stat().st_mode) == 0o700
    assert stat.S_IMODE(service.settings.storage_dir.stat().st_mode) == 0o700
    assert stat.S_IMODE(service.settings.database_path.stat().st_mode) == 0o600
    assert stat.S_IMODE(Path(row["storage_path"]).stat().st_mode) == 0o600


def test_schema_v2_upgrades_to_storage_quota_and_idempotency_state(tmp_path: Path) -> None:
    database = Database(tmp_path / "migration" / "unrender.sqlite3")
    database.initialize()
    with database.connect() as conn:
        conn.execute("DROP TABLE api_idempotency")
        conn.execute("ALTER TABLE jobs DROP COLUMN source_byte_size")
        conn.execute("ALTER TABLE pending_deletions DROP COLUMN user_id")
        conn.execute("ALTER TABLE pending_deletions DROP COLUMN byte_size")
        conn.execute("UPDATE schema_meta SET version=2")

    database.initialize()
    with database.connect() as conn:
        version = conn.execute("SELECT version FROM schema_meta").fetchone()["version"]
        job_columns = {row["name"] for row in conn.execute("PRAGMA table_info(jobs)")}
        deletion_columns = {
            row["name"] for row in conn.execute("PRAGMA table_info(pending_deletions)")
        }
        idempotency = conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name='api_idempotency'"
        ).fetchone()
        assert version == SCHEMA_VERSION
    assert "source_byte_size" in job_columns
    assert {"user_id", "byte_size"} <= deletion_columns
    assert idempotency is not None
    with database.connect() as conn:
        idempotency_columns = {
            row["name"] for row in conn.execute("PRAGMA table_info(api_idempotency)")
        }
    assert {"expires_at", "expired_at"} <= idempotency_columns


@pytest.mark.parametrize("legacy_version", [1, 3])
def test_schema_v1_and_v3_shapes_preserve_rows_foreign_keys_and_expiry(
    tmp_path: Path,
    legacy_version: int,
) -> None:
    database = Database(tmp_path / f"migration-v{legacy_version}" / "unrender.sqlite3")
    database.initialize()
    with database.connect() as conn:
        conn.execute(
            "INSERT INTO users(id,email,password_hash,credit_balance,created_at) "
            "VALUES ('legacy-user','legacy@example.com','hash',0,'2026-01-01T00:00:00Z')"
        )
        conn.execute(
            "INSERT INTO uploads(id,user_id,original_name,mime_type,storage_path,byte_size,"
            "sha256,page_count,created_at,expires_at) VALUES "
            "('legacy-upload','legacy-user','source.png','image/png','/legacy/source.png',321,"
            "'sha',1,'2026-01-01T00:00:00Z','2026-01-02T00:00:00Z')"
        )
        conn.execute(
            "INSERT INTO jobs(id,user_id,upload_id,source_name,source_mime,source_path,"
            "source_sha256,source_byte_size,status,progress_stage,created_at,updated_at) VALUES "
            "('legacy-job','legacy-user','legacy-upload','source.png','image/png',"
            "'/legacy/job.source','sha',321,'queued','Waiting',"
            "'2026-01-01T00:00:00Z','2026-01-01T00:00:00Z')"
        )
        conn.execute("DROP INDEX jobs_lease_expiry_idx")
        conn.execute("DROP INDEX api_idempotency_expiry_idx")
        conn.execute("DROP TABLE provider_attempts")
        conn.execute("DROP TABLE audit_rollups")
        conn.execute("DROP TABLE startup_state")
        for column in (
            "provider_dispatched",
            "provider_dispatched_at",
            "worker_owner",
            "lease_token",
            "lease_generation",
            "lease_expires_at",
            "heartbeat_at",
        ):
            conn.execute(f'ALTER TABLE jobs DROP COLUMN "{column}"')
        if legacy_version == 1:
            conn.execute("DROP TABLE api_idempotency")
            conn.execute("DROP TABLE pending_deletions")
            conn.execute("ALTER TABLE jobs DROP COLUMN source_byte_size")
            conn.execute("ALTER TABLE jobs DROP COLUMN recovery_count")
            conn.execute("ALTER TABLE users DROP COLUMN account_kind")
        else:
            conn.execute("DROP TABLE api_idempotency")
            conn.execute(
                "CREATE TABLE api_idempotency ("
                "user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,"
                "idempotency_key TEXT NOT NULL,request_sha256 TEXT NOT NULL,response_json TEXT,"
                "created_at TEXT NOT NULL,completed_at TEXT,PRIMARY KEY(user_id,idempotency_key))"
            )
            conn.execute(
                "INSERT INTO api_idempotency VALUES "
                "('legacy-user','legacy-key','request-sha','{}',"
                "'2026-01-01T00:00:00Z',NULL)"
            )
            conn.execute("ALTER TABLE pending_deletions DROP COLUMN user_id")
            conn.execute("ALTER TABLE pending_deletions DROP COLUMN byte_size")
        conn.execute("UPDATE schema_meta SET version=?", (legacy_version,))

    database.initialize()
    database.initialize()
    with database.connect() as conn:
        migrated_job = conn.execute(
            "SELECT user_id,upload_id,source_byte_size,lease_generation FROM jobs "
            "WHERE id='legacy-job'"
        ).fetchone()
        migrated_user = conn.execute(
            "SELECT account_kind FROM users WHERE id='legacy-user'"
        ).fetchone()
        assert conn.execute("SELECT version FROM schema_meta").fetchone()[0] == SCHEMA_VERSION
        assert conn.execute("PRAGMA foreign_key_check").fetchone() is None
        assert migrated_job is not None
        assert dict(migrated_job) == {
            "user_id": "legacy-user",
            "upload_id": "legacy-upload",
            "source_byte_size": 321,
            "lease_generation": 0,
        }
        assert migrated_user is not None and migrated_user["account_kind"] == "customer"
        if legacy_version == 3:
            idempotency = conn.execute(
                "SELECT response_json,expires_at,expired_at FROM api_idempotency "
                "WHERE user_id='legacy-user'"
            ).fetchone()
            assert idempotency is not None
            assert idempotency["response_json"] == "{}"
            assert idempotency["expires_at"] == "2026-01-31T00:00:00.000Z"
            assert idempotency["expired_at"] is None


def test_schema_v4_idempotency_migration_preserves_rows_and_enforces_expiry(
    tmp_path: Path,
) -> None:
    database = Database(tmp_path / "migration-v4" / "unrender.sqlite3")
    database.initialize()
    with database.connect() as conn:
        conn.execute("DROP INDEX api_idempotency_expiry_idx")
        conn.execute("DROP TABLE api_idempotency")
        conn.execute(
            """
            CREATE TABLE api_idempotency (
                user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                idempotency_key TEXT NOT NULL,
                request_sha256 TEXT NOT NULL,
                response_json TEXT,
                created_at TEXT NOT NULL,
                completed_at TEXT,
                PRIMARY KEY(user_id, idempotency_key)
            )
            """
        )
        conn.execute(
            """
            INSERT INTO users(id,email,password_hash,credit_balance,created_at)
            VALUES ('legacy-user','legacy@example.com','hash',0,'2026-01-01T00:00:00.000Z')
            """
        )
        conn.execute(
            """
            INSERT INTO api_idempotency(
                user_id,idempotency_key,request_sha256,response_json,created_at,completed_at
            ) VALUES (
                'legacy-user','legacy-key','request-sha','{\"job_id\":\"legacy-job\"}',
                '2026-01-01T00:00:00.000Z','2026-01-01T00:00:01.000Z'
            )
            """
        )
        conn.execute("UPDATE schema_meta SET version=4")

    database.initialize()
    with database.connect() as conn:
        version = conn.execute("SELECT version FROM schema_meta").fetchone()["version"]
        columns = {
            row["name"]: row
            for row in conn.execute("PRAGMA table_info(api_idempotency)").fetchall()
        }
        migrated = conn.execute(
            "SELECT * FROM api_idempotency WHERE user_id='legacy-user'"
        ).fetchone()
        assert version == SCHEMA_VERSION
    assert columns["expires_at"]["notnull"] == 1
    assert migrated["response_json"] == '{"job_id":"legacy-job"}'
    assert migrated["expires_at"] == "2026-01-31T00:00:00.000Z"
    assert migrated["expired_at"] is None


def test_tenant_upload_quota_and_bandwidth_limit_leave_no_orphan_files(
    tmp_path: Path,
) -> None:
    content = png_bytes()
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        max_upload_bytes=len(content) + 100,
        max_user_storage_bytes=(len(content) * 2) + 100,
        max_unattached_uploads=1,
        max_upload_bytes_per_minute=len(content) * 10,
    )
    user_id = customer_id(service)
    first = service.prepare_upload(user_id=user_id, filename="first.png", content=content)
    service.create_job(user_id=user_id, upload_id=first["id"], page_index=0, crop=None)
    with pytest.raises(ProductError, match="storage limit"):
        service.prepare_upload(user_id=user_id, filename="second.png", content=content)
    with service.database.connect() as conn:
        assert conn.execute("SELECT COUNT(*) AS count FROM uploads").fetchone()["count"] == 1
        assert conn.execute("SELECT COUNT(*) AS count FROM jobs").fetchone()["count"] == 1
    assert len(service.storage.object_paths()) == 2

    limited = service_for(
        tmp_path / "bandwidth",
        seed_demo_account=False,
        max_upload_bytes=len(content) + 100,
        max_user_storage_bytes=len(content) * 20,
        max_unattached_uploads=5,
        max_upload_bytes_per_minute=len(content),
    )
    limited_user = customer_id(limited, "bandwidth@example.com")
    limited.prepare_upload(user_id=limited_user, filename="first.png", content=content)
    with pytest.raises(ProductError, match="bandwidth"):
        limited.prepare_upload(user_id=limited_user, filename="second.png", content=content)
    assert len(limited.storage.object_paths()) == 1


def test_saved_sample_runs_free_through_review_approval_and_exports(tmp_path: Path) -> None:
    service = service_for(tmp_path)
    session = service.demo_session()
    user = service.session_user(session["session"])
    assert user is not None
    user_id = str(user["id"])
    starting_credits = service.account(user_id)["credits"]

    upload = service.prepare_demo_upload(user_id)
    job = service.create_job(
        user_id=user_id,
        upload_id=upload["id"],
        page_index=0,
        crop=None,
    )
    assert service.account(user_id)["credits"] == starting_credits
    assert job["status"] == "queued"
    assert service.process_one()

    job = service.get_job(user_id=user_id, job_id=job["id"])
    assert job["status"] == "review"
    assert job["model_version"] == "verified-fixture/synthetic-v1-0002906"
    corrected = job["result"]
    corrected["title"] = "Budget by quarter — reviewed"
    job = service.save_correction(user_id=user_id, job_id=job["id"], result=corrected)
    assert job["result"]["title"].endswith("reviewed")
    versions = service.job_versions(user_id=user_id, job_id=job["id"])
    assert [item["version"] for item in versions["items"]] == [2, 1]
    assert versions["items"][0]["source"] == "correction"
    assert versions["items"][1]["source"] == "extraction"
    assert "result" not in versions["items"][0]
    restored = service.job_version(user_id=user_id, job_id=job["id"], version=1)
    assert restored["result"]["title"] == "Budget (Quarter)"
    job = service.approve(user_id=user_id, job_id=job["id"])
    assert job["status"] == "approved"

    csv_payload, csv_mime = service.export(user_id=user_id, job_id=job["id"], output_format="csv")
    json_payload, _ = service.export(user_id=user_id, job_id=job["id"], output_format="json")
    xlsx_payload, _ = service.export(user_id=user_id, job_id=job["id"], output_format="xlsx")
    assert csv_mime.startswith("text/csv")
    assert b"Q3 2016" in csv_payload
    assert json.loads(json_payload)["title"].endswith("reviewed")
    assert xlsx_payload.startswith(b"PK")
    events = [item["event"] for item in service.job_audit(user_id=user_id, job_id=job["id"])]
    assert "extraction_completed" in events
    assert "result_corrected" in events
    assert "result_approved" in events
    assert events.count("result_exported") == 3


def test_result_history_is_bounded_paginated_and_loaded_one_version_at_a_time(
    tmp_path: Path,
) -> None:
    service = service_for(tmp_path, max_result_versions_per_job=22)
    session = service.demo_session()
    user = service.session_user(session["session"])
    assert user is not None
    user_id = str(user["id"])
    upload = service.prepare_demo_upload(user_id)
    job = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    assert service.process_one()
    for correction in range(21):
        result = service.get_job(user_id=user_id, job_id=job["id"])["result"]
        result["title"] = f"Correction {correction + 1}"
        service.save_correction(user_id=user_id, job_id=job["id"], result=result)

    first_page = service.job_versions(user_id=user_id, job_id=job["id"])
    assert len(first_page["items"]) == 20
    assert first_page["next_before"] == 3
    assert all("result" not in item for item in first_page["items"])
    second_page = service.job_versions(
        user_id=user_id,
        job_id=job["id"],
        before=first_page["next_before"],
    )
    assert [item["version"] for item in second_page["items"]] == [2, 1]
    assert second_page["next_before"] is None
    assert service.job_version(user_id=user_id, job_id=job["id"], version=1)["result"]
    with pytest.raises(ProductError, match="version limit"):
        service.save_correction(
            user_id=user_id,
            job_id=job["id"],
            result=service.get_job(user_id=user_id, job_id=job["id"])["result"],
        )
    with pytest.raises(ProductError, match="version limit"):
        service.reprocess(user_id=user_id, job_id=job["id"])


def test_result_history_byte_quota_rejects_before_charge_or_provider(tmp_path: Path) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        max_history_bytes_per_user=1,
    )
    user_id = customer_id(service)
    upload = service.prepare_upload(
        user_id=user_id,
        filename="verified.webp",
        content=png_bytes(color="navy"),
    )
    with pytest.raises(ProductError, match="maximum-size result"):
        service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    assert service.account(user_id)["credits"] == 3
    assert service.process_one() is False
    with service.database.connect() as conn:
        count = conn.execute("SELECT COUNT(*) AS count FROM result_versions").fetchone()["count"]
        jobs = conn.execute("SELECT COUNT(*) AS count FROM jobs").fetchone()["count"]
        attempts = conn.execute("SELECT COUNT(*) AS count FROM provider_attempts").fetchone()[
            "count"
        ]
    assert count == 0
    assert jobs == 0
    assert attempts == 0


def test_demo_sessions_are_isolated_and_cannot_create_customer_data(tmp_path: Path) -> None:
    app = create_app(settings_for(tmp_path))
    with TestClient(app) as client:
        assert client.post("/api/auth/demo").status_code == 200
        first_account = client.get("/api/me").json()
        assert first_account["demo_account"] is True
        first_upload = client.post("/api/uploads/demo", headers=csrf_headers(client))
        first_job = client.post(
            "/api/jobs",
            headers={**csrf_headers(client), "Idempotency-Key": "demo-first-job"},
            json={"upload_id": first_upload.json()["id"], "page_index": 0},
        )
        assert first_job.status_code == 202
        assert app.state.service.process_one()
        repeated_sample = client.post("/api/uploads/demo", headers=csrf_headers(client))
        assert repeated_sample.status_code == 409

        client.cookies.clear()
        assert client.post("/api/auth/demo").status_code == 200
        second_account = client.get("/api/me").json()
        assert second_account["id"] != first_account["id"]
        assert client.get("/api/jobs").json()["items"] == []
        blocked_upload = client.post(
            "/api/uploads",
            headers=csrf_headers(client),
            files={"file": ("private.png", png_bytes(), "image/png")},
        )
        assert blocked_upload.status_code == 403
        assert (
            client.post("/api/keys", headers=csrf_headers(client), json={"name": "x"}).status_code
            == 403
        )


def test_demo_logout_removes_ephemeral_tenant_and_sources(tmp_path: Path) -> None:
    service = service_for(tmp_path)
    session = service.demo_session()
    user = service.session_user(session["session"])
    assert user is not None
    user_id = str(user["id"])
    upload = service.prepare_demo_upload(user_id)
    service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    assert service.process_one()
    with service.database.connect() as conn:
        paths = [
            Path(row["path"])
            for row in conn.execute(
                "SELECT storage_path AS path FROM uploads WHERE user_id=? "
                "UNION SELECT source_path AS path FROM jobs WHERE user_id=?",
                (user_id, user_id),
            ).fetchall()
        ]
    assert paths and all(path.exists() for path in paths)

    service.logout(session["session"])

    assert service.session_user(session["session"]) is None
    with service.database.connect() as conn:
        assert conn.execute("SELECT 1 FROM users WHERE id=?", (user_id,)).fetchone() is None
        pending = conn.execute("SELECT COUNT(*) AS count FROM pending_deletions").fetchone()
    assert pending["count"] == 0
    assert all(not path.exists() for path in paths)


def test_failed_file_delete_stays_retryable_after_job_row_is_gone(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = service_for(tmp_path, seed_demo_account=False)
    user_id = customer_id(service)
    upload = service.prepare_upload(
        user_id=user_id, filename="delete-me.png", content=png_bytes(color="red")
    )
    job = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    service.cancel(user_id=user_id, job_id=job["id"])
    source_path = Path(service._job_row(user_id=user_id, job_id=job["id"])["source_path"])
    upload_path = Path(service.upload(user_id=user_id, upload_id=upload["id"])["storage_path"])
    original_delete = service.storage.delete

    def fail_delete(_: str | Path) -> None:
        raise OSError("simulated storage outage")

    monkeypatch.setattr(service.storage, "delete", fail_delete)
    assert service.delete_job(user_id=user_id, job_id=job["id"]) is False
    with pytest.raises(ProductError, match="not found"):
        service.get_job(user_id=user_id, job_id=job["id"])
    with pytest.raises(ProductError, match="not found"):
        service.upload_preview(user_id=user_id, upload_id=upload["id"], page_index=0)
    assert source_path.exists() and upload_path.exists()
    with service.database.connect() as conn:
        queued = conn.execute(
            "SELECT storage_path,attempts,last_error FROM pending_deletions "
            "WHERE storage_path IN (?,?) ORDER BY storage_path",
            (str(source_path), str(upload_path)),
        ).fetchall()
        assert conn.execute("SELECT 1 FROM uploads WHERE id=?", (upload["id"],)).fetchone() is None
    assert len(queued) == 2
    assert all(row["attempts"] == 1 and "OSError" in row["last_error"] for row in queued)

    monkeypatch.setattr(service.storage, "delete", original_delete)
    assert service.drain_deletion_queue() == 2
    assert not source_path.exists() and not upload_path.exists()
    with service.database.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM pending_deletions").fetchone()[0] == 0


def test_job_delete_preserves_shared_upload_until_last_reference(tmp_path: Path) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=2)
    user_id = customer_id(service, "shared-upload@example.com")
    upload = service.prepare_upload(
        user_id=user_id, filename="shared.png", content=png_bytes(color="green")
    )
    first = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    second = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    service.cancel(user_id=user_id, job_id=first["id"])
    service.cancel(user_id=user_id, job_id=second["id"])

    assert service.delete_job(user_id=user_id, job_id=first["id"])
    assert service.upload_preview(user_id=user_id, upload_id=upload["id"], page_index=0)
    assert service.delete_job(user_id=user_id, job_id=second["id"])
    with pytest.raises(ProductError, match="not found"):
        service.upload_preview(user_id=user_id, upload_id=upload["id"], page_index=0)


def test_retention_cleanup_queues_files_when_storage_is_temporarily_unavailable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = service_for(tmp_path, seed_demo_account=False)
    user_id = customer_id(service)
    upload = service.prepare_upload(
        user_id=user_id, filename="expired.png", content=png_bytes(color="orange")
    )
    job = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    service.cancel(user_id=user_id, job_id=job["id"])
    with service.database.transaction(immediate=True) as conn:
        paths = [
            str(row["path"])
            for row in conn.execute(
                "SELECT storage_path AS path FROM uploads WHERE id=? "
                "UNION SELECT source_path AS path FROM jobs WHERE id=?",
                (upload["id"], job["id"]),
            ).fetchall()
        ]
        conn.execute(
            "UPDATE uploads SET expires_at='2000-01-01T00:00:00Z' WHERE id=?", (upload["id"],)
        )
        conn.execute("UPDATE jobs SET updated_at='2000-01-01T00:00:00Z' WHERE id=?", (job["id"],))
    original_delete = service.storage.delete

    def fail_retention_delete(_: str | Path) -> None:
        raise OSError("storage unavailable")

    monkeypatch.setattr(service.storage, "delete", fail_retention_delete)

    cleaned = service.cleanup_expired()
    assert cleaned["uploads"] == 1
    assert cleaned["jobs"] == 1
    with service.database.connect() as conn:
        assert conn.execute("SELECT 1 FROM uploads WHERE id=?", (upload["id"],)).fetchone() is None
        assert conn.execute("SELECT 1 FROM jobs WHERE id=?", (job["id"],)).fetchone() is None
        assert (
            conn.execute("SELECT COUNT(*) AS count FROM pending_deletions").fetchone()["count"] >= 2
        )
    assert all(Path(path).exists() for path in paths)

    monkeypatch.setattr(service.storage, "delete", original_delete)
    assert service.drain_deletion_queue() >= 2
    assert all(not Path(path).exists() for path in paths)


def test_storage_reconciliation_removes_unowned_crash_file(tmp_path: Path) -> None:
    service = service_for(tmp_path, seed_demo_account=False, reconciliation_grace_seconds=0)
    orphan = service.settings.storage_dir / "jobs" / "orphan" / "crash.tmp"
    orphan.parent.mkdir(parents=True)
    orphan.write_bytes(b"orphan")
    assert service.reconcile_storage() == 1
    assert service.drain_deletion_queue() == 1
    assert not orphan.exists()


def test_spreadsheet_exports_neutralize_formula_cells(tmp_path: Path) -> None:
    service = service_for(tmp_path)
    session = service.demo_session()
    user = service.session_user(session["session"])
    assert user is not None
    user_id = str(user["id"])
    upload = service.prepare_demo_upload(user_id)
    job = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    assert service.process_one()
    job = service.get_job(user_id=user_id, job_id=job["id"])
    result = job["result"]
    result["x_axis"]["label"] = "=SUM(A1:A2)"
    result["y_axis"]["label"] = "@external"
    result["series"][0]["points"][0] = {
        "x": "+cmd|' /C calc'!A0",
        "y": -9.2,
    }
    service.save_correction(user_id=user_id, job_id=job["id"], result=result)

    csv_payload, _ = service.export(user_id=user_id, job_id=job["id"], output_format="csv")
    rows = list(csv.reader(io.StringIO(csv_payload.decode("utf-8"))))
    assert rows[0] == ["'=SUM(A1:A2)", "'@external"]
    assert rows[1] == ["'+cmd|' /C calc'!A0", "-9.2"]

    xlsx_payload, _ = service.export(user_id=user_id, job_id=job["id"], output_format="xlsx")
    workbook = load_workbook(io.BytesIO(xlsx_payload), data_only=False)
    sheet = workbook["Extracted data"]
    assert sheet["A1"].value == "'=SUM(A1:A2)"
    assert sheet["A1"].data_type == "s"
    assert sheet["A2"].value == "'+cmd|' /C calc'!A0"
    assert sheet["B2"].value == -9.2
    assert sheet["B2"].data_type == "n"


def test_multi_series_exports_preserve_duplicate_points_and_numeric_cells(tmp_path: Path) -> None:
    service = service_for(tmp_path)
    chart = ChartData.model_validate(
        {
            "chart_type": "multi_line",
            "title": "Repeated categories",
            "x_axis": {"label": "Period", "unit": None},
            "y_axis": {"label": "Value", "unit": None},
            "series": [
                {"name": None, "points": [{"x": "001", "y": 1.25}, {"x": "001", "y": 2.5}]},
                {"name": "Peer", "points": [{"x": 1.0, "y": 3.75}]},
            ],
        }
    )
    csv_rows = list(csv.reader(io.StringIO(service._safe_csv(chart))))
    assert csv_rows == [
        ["series", "Period", "Value"],
        ["series_1", "001", "1.25"],
        ["series_1", "001", "2.5"],
        ["Peer", "1.0", "3.75"],
    ]
    workbook = load_workbook(
        io.BytesIO(
            service._xlsx(
                chart,
                {
                    "source_name": "multi.png",
                    "id": "job-1",
                    "model_version": "test-model",
                    "status": "review",
                    "approved_at": None,
                },  # type: ignore[arg-type]
            )
        )
    )
    sheet = workbook["Extracted data"]
    assert sheet["B2"].value == "001"
    assert sheet["C2"].value == 1.25
    assert sheet["C2"].data_type == "n"


def test_failed_and_cancelled_jobs_refund_once_and_tenants_are_isolated(
    tmp_path: Path,
) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=2)
    first = service.register("first@example.com", "first password is long enough")
    second = service.register("second@example.com", "second password is long enough")
    first_user = service.session_user(first["session"])
    second_user = service.session_user(second["session"])
    assert first_user is not None and second_user is not None
    first_id, second_id = str(first_user["id"]), str(second_user["id"])

    upload = service.prepare_upload(
        user_id=first_id, filename="custom.png", content=png_bytes(color="navy")
    )
    failed = service.create_job(user_id=first_id, upload_id=upload["id"], page_index=0, crop=None)
    assert service.account(first_id)["credits"] == 1
    assert service.process_one()
    failed = service.get_job(user_id=first_id, job_id=failed["id"])
    assert failed["status"] == "failed"
    assert failed["error"]["code"] == "provider_not_configured"
    assert service.account(first_id)["credits"] == 1
    with pytest.raises(ProductError, match="not found"):
        service.get_job(user_id=second_id, job_id=failed["id"])

    queued = service.create_job(user_id=first_id, upload_id=upload["id"], page_index=0, crop=None)
    assert service.account(first_id)["credits"] == 0
    cancelled = service.cancel(user_id=first_id, job_id=queued["id"])
    assert cancelled["status"] == "cancelled"
    assert service.account(first_id)["credits"] == 1
    with pytest.raises(ProductError, match="cannot be cancelled"):
        service.cancel(user_id=first_id, job_id=queued["id"])
    assert service.account(first_id)["credits"] == 1


def test_interrupted_job_has_one_automatic_recovery_then_refunds(tmp_path: Path) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=1)
    session = service.register("recover@example.com", "recovery password is long enough")
    user = service.session_user(session["session"])
    assert user is not None
    user_id = str(user["id"])
    upload = service.prepare_upload(
        user_id=user_id, filename="custom.png", content=png_bytes(color="black")
    )
    job = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    claimed = service.claim_next_job()
    assert claimed is not None and claimed["status"] == "running"
    with service.database.transaction(immediate=True) as conn:
        conn.execute(
            "UPDATE jobs SET lease_expires_at=? WHERE id=?",
            (timestamp(utcnow() - timedelta(seconds=1)), job["id"]),
        )
    assert service.recover_interrupted_jobs() == 1
    recovered = service.get_job(user_id=user_id, job_id=job["id"])
    assert recovered["status"] == "queued"
    assert recovered["recovery_count"] == 1
    assert service.account(user_id)["credits"] == 0
    assert service.claim_next_job() is not None
    with service.database.transaction(immediate=True) as conn:
        conn.execute(
            "UPDATE jobs SET lease_expires_at=? WHERE id=?",
            (timestamp(utcnow() - timedelta(seconds=1)), job["id"]),
        )
    assert service.recover_interrupted_jobs() == 1
    exhausted = service.get_job(user_id=user_id, job_id=job["id"])
    assert exhausted["status"] == "failed"
    assert exhausted["error"]["code"] == "worker_recovery_exhausted"
    assert service.account(user_id)["credits"] == 1


def test_billing_event_is_idempotent_and_detects_conflicting_replay(tmp_path: Path) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        initial_credits=0,
        stripe_secret_key="sk_test_local",
        stripe_webhook_secret="whsec_local",
        stripe_price_id="price_local",
        max_billing_events_global=1,
    )
    session = service.register("billing@example.com", "billing password is long enough")
    user = service.session_user(session["session"])
    assert user is not None
    user_id = str(user["id"])
    arguments = {
        "event_id": "evt_test_1",
        "event_type": "checkout.session.completed",
        "user_id": user_id,
        "credits": 100,
        "payload_sha256": "abc123",
    }
    assert service.apply_billing_event(**arguments)
    assert not service.apply_billing_event(**arguments)
    assert service.account(user_id)["credits"] == 100
    with pytest.raises(ProductError, match="did not match"):
        service.apply_billing_event(**{**arguments, "payload_sha256": "different"})
    assert service.account(user_id)["credits"] == 100

    assert not service.apply_billing_event(**arguments)
    with pytest.raises(ProductError, match="processing is paused"):
        service.apply_billing_event(**{**arguments, "event_id": "evt_test_2"})


def test_http_flow_enforces_csrf_origin_headers_and_api_tenancy(tmp_path: Path) -> None:
    app = create_app(settings_for(tmp_path))
    with TestClient(app) as client:
        short_password = client.post(
            "/api/auth/register",
            json={"email": "short@example.com", "password": "too-short"},
        )
        assert short_password.status_code == 422
        oversized = client.post(
            "/api/auth/login",
            headers={"Content-Length": str(22 * 1024 * 1024)},
            content=b"{}",
        )
        assert oversized.status_code == 413
        assert oversized.headers["content-security-policy"].startswith("default-src 'self'")
        assert client.get("/", headers={"Host": "attacker.example"}).status_code == 400
        unexpected = client.post(
            "/api/auth/register",
            json={
                "email": "extra@example.com",
                "password": "a sufficiently long password",
                "admin": True,
            },
        )
        assert unexpected.status_code == 422
        response = client.post("/api/auth/demo")
        assert response.status_code == 200
        assert client.get("/").headers["x-frame-options"] == "DENY"
        assert client.post("/api/uploads/demo").status_code == 403
        assert (
            client.post(
                "/api/uploads/demo",
                headers={**csrf_headers(client), "Origin": "https://attacker.example"},
            ).status_code
            == 403
        )

        upload = client.post("/api/uploads/demo", headers=csrf_headers(client))
        assert upload.status_code == 201
        job = client.post(
            "/api/jobs",
            headers={**csrf_headers(client), "Idempotency-Key": "demo-http-job"},
            json={"upload_id": upload.json()["id"], "page_index": 0},
        )
        assert job.status_code == 202
        assert app.state.service.process_one()
        job_id = job.json()["id"]
        result = client.get(f"/api/jobs/{job_id}").json()
        assert result["status"] == "review"
        assert client.get("/api/me").json()["demo_account"] is True
        assert (
            client.post(
                "/api/uploads",
                headers=csrf_headers(client),
                files={"file": ("private.png", png_bytes(), "image/png")},
            ).status_code
            == 403
        )
        assert (
            client.post(
                "/api/keys",
                headers=csrf_headers(client),
                json={"name": "blocked demo key"},
            ).status_code
            == 403
        )
        assert client.post("/api/billing/checkout", headers=csrf_headers(client)).status_code == 403
        assert client.post("/api/auth/logout", headers=csrf_headers(client)).status_code == 200

        registered = client.post(
            "/api/auth/register",
            json={"email": "integration@example.com", "password": "integration password is long"},
        )
        assert registered.status_code == 201
        customer_upload = client.post(
            "/api/uploads",
            headers=csrf_headers(client),
            files={"file": ("customer.png", png_bytes(color="green"), "image/png")},
        )
        assert customer_upload.status_code == 201
        account_response = client.get("/api/me")
        assert account_response.status_code == 200
        assert len(account_response.json()["principal_marker"]) == 32
        missing_key = client.post(
            "/api/jobs",
            headers=csrf_headers(client),
            json={"upload_id": customer_upload.json()["id"], "page_index": 0},
        )
        assert missing_key.status_code == 422
        assert missing_key.json()["error"]["code"] == "idempotency_key_required"
        assert client.get("/api/jobs").json()["items"] == []
        customer_job = client.post(
            "/api/jobs",
            headers={**csrf_headers(client), "Idempotency-Key": "customer-http-job"},
            json={"upload_id": customer_upload.json()["id"], "page_index": 0},
        )
        assert customer_job.status_code == 202
        customer_job_id = customer_job.json()["id"]
        assert client.get(f"/api/jobs/{job_id}").status_code == 404
        key_response = client.post(
            "/api/keys",
            headers=csrf_headers(client),
            json={"name": "test integration"},
        )
        assert key_response.status_code == 201
        secret = key_response.json()["key"]
        api_result = client.get(
            f"/api/v1/extractions/{customer_job_id}",
            headers={"Authorization": f"Bearer {secret}"},
        )
        assert api_result.status_code == 200
        assert api_result.json()["id"] == customer_job_id
        keys = client.get("/api/keys")
        assert keys.status_code == 200
        key_id = keys.json()["items"][0]["id"]
        assert "key_hash" not in keys.json()["items"][0]
        revoked = client.delete(f"/api/keys/{key_id}", headers=csrf_headers(client))
        assert revoked.status_code == 200
        assert revoked.json()["revoked_at"]
        assert (
            client.get(
                f"/api/v1/extractions/{customer_job_id}",
                headers={"Authorization": f"Bearer {secret}"},
            ).status_code
            == 401
        )


def test_api_submission_is_tenant_idempotent_and_conflicts_on_request_drift(
    tmp_path: Path,
) -> None:
    app = create_app(settings_for(tmp_path, seed_demo_account=False))
    with TestClient(app) as client:
        secret = register_and_create_api_key(client)
        headers = {
            "Authorization": f"Bearer {secret}",
            "Idempotency-Key": "extract-2026-0001",
        }
        source = png_bytes(color="green")
        first = client.post(
            "/api/v1/extractions?page_index=0",
            headers=headers,
            files={"file": ("chart.png", source, "image/png")},
        )
        replay = client.post(
            "/api/v1/extractions?page_index=0",
            headers=headers,
            files={"file": ("chart.png", source, "image/png")},
        )
        assert first.status_code == replay.status_code == 202
        assert first.json() == replay.json()
        assert client.get("/api/me").json()["credits"] == 2

        conflict = client.post(
            "/api/v1/extractions?page_index=0",
            headers=headers,
            files={"file": ("different.png", png_bytes(color="red"), "image/png")},
        )
        assert conflict.status_code == 409
        assert conflict.json()["error"]["code"] == "idempotency_conflict"
        missing = client.post(
            "/api/v1/extractions?page_index=0",
            headers={"Authorization": f"Bearer {secret}"},
            files={"file": ("chart.png", source, "image/png")},
        )
        assert missing.status_code == 422

        with app.state.service.database.connect() as conn:
            assert conn.execute("SELECT COUNT(*) AS count FROM uploads").fetchone()["count"] == 1
            assert conn.execute("SELECT COUNT(*) AS count FROM jobs").fetchone()["count"] == 1
            assert (
                conn.execute("SELECT COUNT(*) AS count FROM api_idempotency").fetchone()["count"]
                == 1
            )


def test_expired_idempotency_key_is_tombstoned_through_the_retention_window(
    tmp_path: Path,
) -> None:
    app = create_app(
        settings_for(
            tmp_path,
            seed_demo_account=False,
            idempotency_ttl_hours=1,
        )
    )
    with TestClient(app) as client:
        secret = register_and_create_api_key(client, email="expiry@example.com")
        headers = {
            "Authorization": f"Bearer {secret}",
            "Idempotency-Key": "expiry-window-0001",
        }
        source = png_bytes(color="green")
        first = client.post(
            "/api/v1/extractions",
            headers=headers,
            files={"file": ("chart.png", source, "image/png")},
        )
        assert first.status_code == 202
        with app.state.service.database.transaction(immediate=True) as conn:
            conn.execute(
                "UPDATE api_idempotency SET expires_at=?",
                (timestamp(utcnow() - timedelta(minutes=1)),),
            )
        app.state.service.cleanup_expired()

        replay = client.post(
            "/api/v1/extractions",
            headers=headers,
            files={"file": ("chart.png", source, "image/png")},
        )
        assert replay.status_code == 409
        assert replay.json()["error"]["code"] == "idempotency_key_expired"
        assert client.get("/api/me").json()["credits"] == 2
        with app.state.service.database.connect() as conn:
            assert conn.execute("SELECT COUNT(*) AS count FROM jobs").fetchone()["count"] == 1
            tombstone = conn.execute(
                "SELECT response_json,expired_at FROM api_idempotency"
            ).fetchone()
        assert tombstone["response_json"] is None
        assert tombstone["expired_at"] is not None

        with app.state.service.database.transaction(immediate=True) as conn:
            conn.execute(
                "UPDATE api_idempotency SET expired_at=?",
                (timestamp(utcnow() - timedelta(days=366)),),
            )
        app.state.service.cleanup_expired()
        with app.state.service.database.connect() as conn:
            assert (
                conn.execute("SELECT COUNT(*) AS count FROM api_idempotency").fetchone()["count"]
                == 0
            )


def test_zero_credit_api_submission_leaves_no_rows_or_files(tmp_path: Path) -> None:
    app = create_app(
        settings_for(
            tmp_path,
            seed_demo_account=False,
            initial_credits=0,
        )
    )
    with TestClient(app) as client:
        secret = register_and_create_api_key(client, email="zero@example.com")
        rejected = client.post(
            "/api/v1/extractions",
            headers={
                "Authorization": f"Bearer {secret}",
                "Idempotency-Key": "zero-credit-0001",
            },
            files={"file": ("chart.png", png_bytes(), "image/png")},
        )
        assert rejected.status_code == 402
        with app.state.service.database.connect() as conn:
            assert conn.execute("SELECT COUNT(*) AS count FROM uploads").fetchone()["count"] == 0
            assert conn.execute("SELECT COUNT(*) AS count FROM jobs").fetchone()["count"] == 0
            assert (
                conn.execute("SELECT COUNT(*) AS count FROM api_idempotency").fetchone()["count"]
                == 0
            )
        assert app.state.service.storage.object_paths() == []


def test_streaming_body_limit_rejects_chunked_payload_before_json_parsing(
    tmp_path: Path,
) -> None:
    app = create_app(settings_for(tmp_path, seed_demo_account=False))

    def oversized_body() -> Iterator[bytes]:
        for _ in range(300):
            yield b"x" * 1024

    with TestClient(app) as client:
        response = client.post(
            "/api/auth/login",
            headers={"Content-Type": "application/json"},
            content=oversized_body(),
        )
    assert response.status_code == 413
    assert response.json()["error"]["code"] == "request_too_large"
    assert "default-src 'self'" in response.headers["Content-Security-Policy"]

    consumed = 0

    def hostile_host_body() -> Iterator[bytes]:
        nonlocal consumed
        for _ in range(300):
            consumed += 1
            yield b"x" * 1024

    with TestClient(app) as client:
        hostile = client.post(
            "/api/auth/login",
            headers={"Content-Type": "application/json", "Host": "attacker.example"},
            content=hostile_host_body(),
        )
    assert hostile.status_code == 400
    assert consumed == 0
    assert "default-src 'self'" in hostile.headers["Content-Security-Policy"]


def test_body_limit_backpressures_concurrent_chunks_without_aggregate_buffering() -> None:
    stream_count = 32
    consumed = [0] * stream_count
    statuses: list[int] = []

    async def consume_app(scope: Any, receive: Any, send: Any) -> None:
        del scope
        while True:
            message = await receive()
            if message["type"] == "http.disconnect":
                return
            if not message.get("more_body", False):
                await send({"type": "http.response.start", "status": 204, "headers": []})
                await send({"type": "http.response.body", "body": b""})
                return

    middleware = BodyLimitMiddleware(consume_app, upload_limit=20 * 1024 * 1024)

    async def invoke(index: int) -> None:
        async def receive() -> dict[str, object]:
            consumed[index] += 1
            await asyncio.sleep(0)
            return {"type": "http.request", "body": b"x" * 65536, "more_body": True}

        async def send(message: dict[str, object]) -> None:
            if message["type"] == "http.response.start":
                statuses.append(int(message["status"]))

        await middleware({"type": "http", "path": "/api/auth/login"}, receive, send)

    async def exercise() -> None:
        await asyncio.gather(*(invoke(index) for index in range(stream_count)))

    tracemalloc.start()
    try:
        asyncio.run(exercise())
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert statuses == [413] * stream_count
    assert consumed == [5] * stream_count
    assert peak < 16 * 1024 * 1024


def test_test_mode_checkout_and_signed_webhook_flow(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    app = create_app(
        settings_for(
            tmp_path,
            seed_demo_account=False,
            initial_credits=0,
            stripe_secret_key="sk_test_local",
            stripe_webhook_secret="whsec_local",
            stripe_price_id="price_local",
        )
    )
    checkout_arguments: dict[str, object] = {}

    def fake_checkout(**kwargs: object) -> SimpleNamespace:
        checkout_arguments.update(kwargs)
        return SimpleNamespace(url="https://checkout.stripe.com/c/pay/cs_test_local")

    monkeypatch.setattr(stripe.checkout.Session, "create", fake_checkout)
    with TestClient(app) as client:
        registered = client.post(
            "/api/auth/register",
            json={"email": "buyer@example.com", "password": "buyer password is long enough"},
        )
        assert registered.status_code == 201
        account = client.get("/api/me").json()
        checkout = client.post("/api/billing/checkout", headers=csrf_headers(client))
        assert checkout.status_code == 201
        assert checkout.json()["url"].startswith("https://checkout.stripe.com/")
        assert checkout_arguments["payment_method_types"] == ["card"]

        monkeypatch.setattr(
            stripe.checkout.Session,
            "create",
            lambda **_: SimpleNamespace(url="https://attacker.example/cs_test_local"),
        )
        rejected_checkout = client.post("/api/billing/checkout", headers=csrf_headers(client))
        assert rejected_checkout.status_code == 502

        event = {
            "id": "evt_signed_test",
            "type": "checkout.session.completed",
            "data": {
                "object": {
                    "payment_status": "paid",
                    "client_reference_id": account["id"],
                    "metadata": {"user_id": account["id"], "credits": "100"},
                }
            },
        }
        monkeypatch.setattr(stripe.Webhook, "construct_event", lambda *_: event)
        first = client.post(
            "/api/billing/webhook",
            content=b"signed-payload",
            headers={"Stripe-Signature": "test-signature"},
        )
        second = client.post(
            "/api/billing/webhook",
            content=b"signed-payload",
            headers={"Stripe-Signature": "test-signature"},
        )
        assert first.json() == {"received": True, "applied": True}
        assert second.json() == {"received": True, "applied": False}
        assert client.get("/api/me").json()["credits"] == 100


def test_rate_limit_cannot_be_bypassed_with_rotated_credentials_or_forwarded_ip(
    tmp_path: Path,
) -> None:
    app = create_app(
        settings_for(
            tmp_path,
            seed_demo_account=False,
            rate_limit_per_minute=2,
            auth_rate_limit_per_minute=2,
        )
    )
    payload = {"email": "known@example.com", "password": "wrong password"}
    with TestClient(app) as client:
        app.state.service.register("known@example.com", "known password is long enough")
        first = client.post(
            "/api/auth/login",
            json=payload,
            headers={"Authorization": "Bearer attacker-a", "X-Forwarded-For": "198.51.100.1"},
        )
        second = client.post(
            "/api/auth/login",
            json=payload,
            headers={
                "Cookie": "unrender_session=attacker-b",
                "X-Forwarded-For": "198.51.100.2",
            },
        )
        blocked = client.post(
            "/api/auth/login",
            json=payload,
            headers={"Authorization": "Bearer attacker-c", "X-Forwarded-For": "198.51.100.3"},
        )
    assert first.status_code == second.status_code == 401
    assert blocked.status_code == 429
    with app.state.service.database.connect() as conn:
        buckets = conn.execute("SELECT bucket_key,request_count FROM rate_limits").fetchall()
    assert len(buckets) == 1
    assert int(buckets[0]["request_count"]) == 3
    assert "attacker" not in str(buckets[0]["bucket_key"])


def test_live_health_is_cheap_while_readiness_is_admission_limited(tmp_path: Path) -> None:
    app = create_app(settings_for(tmp_path, rate_limit_per_minute=2))
    with TestClient(app) as client:
        assert client.get("/health/ready").status_code == 200
        assert client.get("/health/ready").status_code == 200
        assert client.get("/health/ready").status_code == 429
        assert client.get("/health/live").status_code == 200


@pytest.mark.parametrize("filename", ["index.html", "privacy.html", "terms.html"])
def test_static_pages_keep_basic_accessibility_contract(filename: str) -> None:
    soup = BeautifulSoup((STATIC_DIR / filename).read_text(encoding="utf-8"), "html.parser")
    ids = [tag["id"] for tag in soup.find_all(attrs={"id": True})]
    assert len(ids) == len(set(ids))
    assert soup.html is not None and soup.html.get("lang") == "en"
    assert soup.title is not None and soup.title.get_text(strip=True)
    assert soup.find("main") is not None
    assert soup.find("h1") is not None
    for image in soup.find_all("img"):
        assert image.get("alt") is not None
    for button in soup.find_all("button"):
        assert button.get("type") in {"button", "submit", "reset"}
    for field in soup.find_all(["input", "select"]):
        if field.get("type") == "hidden":
            continue
        field_id = field.get("id")
        wrapped = field.find_parent("label") is not None
        referenced = bool(field_id and soup.find("label", attrs={"for": field_id}))
        assert wrapped or referenced, f"Unlabelled field: {field}"
    for dialog in soup.find_all("dialog"):
        label_id = dialog.get("aria-labelledby")
        assert label_id and soup.find(id=label_id)


def _paid_job(service: ProductService, user_id: str, *, color: str = "white") -> dict[str, object]:
    upload = service.prepare_upload(
        user_id=user_id,
        filename=f"{color}.png",
        content=png_bytes(color=color),
    )
    return service.create_job(
        user_id=user_id,
        upload_id=str(upload["id"]),
        page_index=0,
        crop=None,
    )


def _install_successful_extractor(service: ProductService) -> None:
    fixture = json.loads(
        (STATIC_DIR / "demo" / "budget-quarter-result.json").read_text(encoding="utf-8")
    )
    service.extractor = SimpleNamespace(
        extract=lambda _: ExtractionOutput(
            chart=ChartData.model_validate(fixture["result"]),
            raw="deterministic test output",
            extractor="test",
            model_version="test-pinned",
        )
    )


def test_worker_leases_are_fresh_reaped_once_and_fence_every_terminal_mutation(
    tmp_path: Path,
) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        initial_credits=1,
        reconciliation_grace_seconds=0,
    )
    user_id = customer_id(service, "lease@example.com")
    job = _paid_job(service, user_id, color="navy")
    stale_claim = service.claim_next_job("service-a")
    assert stale_claim is not None

    peer = ProductService(
        settings=service.settings,
        database=Database(service.settings.database_path),
        storage=Storage(service.settings),
        extractor=ReplayExtractor(STATIC_DIR),
        static_dir=STATIC_DIR,
    )
    peer.initialize()
    with service.database.connect() as conn:
        fresh = conn.execute(
            "SELECT status,worker_owner,lease_token,source_path FROM jobs WHERE id=?", (job["id"],)
        ).fetchone()
        schema_version = conn.execute("SELECT version FROM schema_meta").fetchone()[0]
        started_events = conn.execute(
            "SELECT COUNT(*) FROM audit_events WHERE job_id=? AND event_type='job_started'",
            (job["id"],),
        ).fetchone()[0]
        provider_attempts = conn.execute(
            "SELECT COUNT(*) FROM provider_attempts WHERE job_id=?", (job["id"],)
        ).fetchone()[0]
    assert fresh["status"] == "running"
    assert fresh["worker_owner"] == "service-a"
    assert fresh["lease_token"] == stale_claim.token
    assert Path(fresh["source_path"]).is_file()
    assert schema_version == SCHEMA_VERSION
    assert started_events == 1
    assert provider_attempts == 0

    with service.database.transaction(immediate=True) as conn:
        conn.execute(
            "UPDATE jobs SET lease_expires_at=? WHERE id=?",
            (timestamp(utcnow() - timedelta(seconds=1)), job["id"]),
        )
    barrier = threading.Barrier(3)
    reaped: list[int] = []

    def recover(candidate: ProductService) -> None:
        barrier.wait()
        reaped.append(candidate.recover_interrupted_jobs())

    threads = [threading.Thread(target=recover, args=(candidate,)) for candidate in (service, peer)]
    for thread in threads:
        thread.start()
    barrier.wait()
    for thread in threads:
        thread.join(timeout=5)
    assert all(not thread.is_alive() for thread in threads)
    assert sorted(reaped) == [0, 1]

    current_claim = peer.claim_next_job("service-b")
    assert current_claim is not None
    assert current_claim.generation == stale_claim.generation + 1
    assert not service.heartbeat_claim(stale_claim)
    assert not service._update_claim_progress(stale_claim, "stale progress")
    assert not service._finish_failed_claim(stale_claim, "stale", "must not commit")
    assert not service._finish_cancelled_claim(stale_claim)
    assert service.account(user_id)["credits"] == 0
    assert peer._finish_failed_claim(current_claim, "pre_dispatch", "safe failure")
    assert not peer._finish_failed_claim(current_claim, "duplicate", "must not commit")
    assert service.account(user_id)["credits"] == 1
    with service.database.connect() as conn:
        stale_events = conn.execute(
            "SELECT COUNT(*) AS count FROM audit_events "
            "WHERE job_id=? AND details_json LIKE '%stale%'",
            (job["id"],),
        ).fetchone()["count"]
        refunds = conn.execute(
            "SELECT COUNT(*) AS count FROM credit_ledger WHERE idempotency_key=?",
            (f"job:{job['id']}:refund:1",),
        ).fetchone()["count"]
    assert stale_events == 0
    assert refunds == 1


def test_predispatch_cancel_avoids_provider_and_postdispatch_failures_open_circuit(
    tmp_path: Path,
) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        initial_credits=4,
        provider_failure_limit_per_user_hour=1,
        provider_failure_limit_global_hour=1,
    )
    user_id = customer_id(service, "circuit@example.com")
    cancelled = _paid_job(service, user_id, color="blue")
    claim = service.claim_next_job("cancel-owner")
    assert claim is not None
    service.cancel(user_id=user_id, job_id=str(cancelled["id"]))
    assert not service._begin_provider_dispatch(claim)
    assert service.get_job(user_id=user_id, job_id=str(cancelled["id"]))["status"] == "cancelled"
    assert service.account(user_id)["credits"] == 4

    charged_cancel = _paid_job(service, user_id, color="black")
    charged_claim = service.claim_next_job("charged-cancel-owner")
    assert charged_claim is not None and service._begin_provider_dispatch(charged_claim)
    service.cancel(user_id=user_id, job_id=str(charged_cancel["id"]))
    assert service._finish_cancelled_claim(charged_claim)
    assert service.get_job(user_id=user_id, job_id=str(charged_cancel["id"]))["status"] == (
        "cancelled"
    )
    assert service.account(user_id)["credits"] == 3

    provider_calls = 0

    def fail_provider(_: bytes) -> None:
        nonlocal provider_calls
        provider_calls += 1
        raise ExtractionError("provider_unavailable", "provider failed")

    service.extractor = SimpleNamespace(extract=fail_provider)
    first = _paid_job(service, user_id, color="red")
    assert service.process_one("provider-owner")
    assert service.get_job(user_id=user_id, job_id=str(first["id"]))["status"] == "failed"
    assert service.account(user_id)["credits"] == 2

    second = _paid_job(service, user_id, color="green")
    assert service.account(user_id)["credits"] == 1
    assert service.process_one("provider-owner")
    blocked = service.get_job(user_id=user_id, job_id=str(second["id"]))
    assert blocked["status"] == "failed"
    assert blocked["error"]["code"] == "provider_circuit_open"
    assert provider_calls == 1
    assert service.account(user_id)["credits"] == 2


def test_provider_attempt_history_is_bounded_before_another_dispatch(tmp_path: Path) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        initial_credits=3,
        max_provider_attempts_per_job=1,
        provider_failure_limit_per_user_hour=10,
        provider_failure_limit_global_hour=10,
    )
    user_id = customer_id(service, "attempt-cap@example.com")
    job = _paid_job(service, user_id, color="purple")
    first = service.claim_next_job("attempt-owner-1")
    assert first is not None and service._begin_provider_dispatch(first)
    assert service._finish_failed_claim(first, "provider_failed", "charged failure")
    assert service.account(user_id)["credits"] == 2

    service.reprocess(user_id=user_id, job_id=str(job["id"]))
    second = service.claim_next_job("attempt-owner-2")
    assert second is not None
    assert not service._begin_provider_dispatch(second)
    assert service.account(user_id)["credits"] == 2
    current = service.get_job(user_id=user_id, job_id=str(job["id"]))
    assert current["status"] == "failed"
    assert current["error"]["code"] == "provider_attempt_history_limit"
    with service.database.connect() as conn:
        assert (
            conn.execute(
                "SELECT COUNT(*) FROM provider_attempts WHERE job_id=?", (job["id"],)
            ).fetchone()[0]
            == 1
        )


def test_worker_shutdown_keeps_heartbeat_until_long_provider_call_finishes(
    tmp_path: Path,
) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        initial_credits=1,
        worker_lease_seconds=10,
        worker_heartbeat_seconds=3,
        worker_shutdown_timeout_seconds=1,
    )
    user_id = customer_id(service, "drain@example.com")
    job = _paid_job(service, user_id, color="purple")
    fixture = json.loads(
        (STATIC_DIR / "demo" / "budget-quarter-result.json").read_text(encoding="utf-8")
    )
    provider_started = threading.Event()
    provider_release = threading.Event()

    def slow_provider(_: bytes) -> ExtractionOutput:
        provider_started.set()
        assert provider_release.wait(timeout=10)
        return ExtractionOutput(
            chart=ChartData.model_validate(fixture["result"]),
            raw="slow provider",
            extractor="test",
            model_version="pinned",
        )

    service.extractor = SimpleNamespace(extract=slow_provider)
    worker = JobWorker(service, poll_seconds=0.01)
    worker.start()
    assert provider_started.wait(timeout=5)
    with service.database.connect() as conn:
        before = conn.execute(
            "SELECT lease_expires_at FROM jobs WHERE id=?", (job["id"],)
        ).fetchone()["lease_expires_at"]
    assert not worker.stop(timeout=0.05)
    time.sleep(3.5)
    peer = ProductService(
        settings=service.settings,
        database=Database(service.settings.database_path),
        storage=Storage(service.settings),
        extractor=ReplayExtractor(STATIC_DIR),
        static_dir=STATIC_DIR,
    )
    peer.initialize()
    assert peer.recover_interrupted_jobs() == 0
    with service.database.connect() as conn:
        after = conn.execute(
            "SELECT lease_expires_at,status FROM jobs WHERE id=?", (job["id"],)
        ).fetchone()
    assert after["lease_expires_at"] > before
    assert after["status"] == "running"
    provider_release.set()
    assert worker.stop(timeout=5)
    assert service.get_job(user_id=user_id, job_id=str(job["id"]))["status"] == "review"
    assert service.account(user_id)["credits"] == 0


def test_migration_failure_is_atomic_and_concurrent_initializers_converge(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    database = Database(tmp_path / "migration-atomic" / "unrender.sqlite3")
    database.initialize()
    with database.connect() as conn:
        conn.execute("DROP INDEX api_idempotency_expiry_idx")
        conn.execute("ALTER TABLE api_idempotency RENAME TO api_idempotency_v4")
        conn.execute(
            "CREATE TABLE api_idempotency ("
            "user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,"
            "idempotency_key TEXT NOT NULL,request_sha256 TEXT NOT NULL,response_json TEXT,"
            "created_at TEXT NOT NULL,completed_at TEXT,PRIMARY KEY(user_id,idempotency_key))"
        )
        conn.execute(
            "INSERT INTO users(id,email,password_hash,credit_balance,created_at) "
            "VALUES ('atomic-user','atomic@example.com','hash',0,'2026-01-01T00:00:00Z')"
        )
        conn.execute(
            "INSERT INTO api_idempotency VALUES "
            "('atomic-user','key','sha','{}','2026-01-01T00:00:00Z',NULL)"
        )
        conn.execute("DROP TABLE api_idempotency_v4")
        conn.execute("UPDATE schema_meta SET version=4")

    original = database._migrate_v4_to_v5

    def crash_after_migration(conn: sqlite3.Connection) -> None:
        original(conn)
        raise RuntimeError("injected migration crash")

    monkeypatch.setattr(database, "_migrate_v4_to_v5", crash_after_migration)
    with pytest.raises(RuntimeError, match="injected"):
        database.initialize()
    with database.connect() as conn:
        assert conn.execute("SELECT version FROM schema_meta").fetchone()["version"] == 4
        assert "expires_at" not in {
            row["name"] for row in conn.execute("PRAGMA table_info(api_idempotency)")
        }
        assert conn.execute("SELECT response_json FROM api_idempotency").fetchone()[0] == "{}"

    monkeypatch.setattr(database, "_migrate_v4_to_v5", original)
    errors: list[Exception] = []
    barrier = threading.Barrier(7)

    def initialize_peer() -> None:
        barrier.wait()
        try:
            Database(database.path).initialize()
        except Exception as exc:  # pragma: no cover - assertion reports the exact peer error
            errors.append(exc)

    threads = [threading.Thread(target=initialize_peer) for _ in range(6)]
    for thread in threads:
        thread.start()
    barrier.wait()
    for thread in threads:
        thread.join(timeout=10)
    assert all(not thread.is_alive() for thread in threads)
    assert errors == []
    database.initialize()
    with database.connect() as conn:
        assert (
            conn.execute("SELECT version FROM schema_meta").fetchone()["version"] == SCHEMA_VERSION
        )
        assert conn.execute("SELECT response_json FROM api_idempotency").fetchone()[0] == "{}"
        assert conn.execute("PRAGMA foreign_key_check").fetchone() is None


@pytest.mark.parametrize("crash_after", range(1, 7))
def test_v4_to_v5_rolls_back_after_each_mutating_statement(
    tmp_path: Path, crash_after: int
) -> None:
    database = Database(tmp_path / f"migration-crash-{crash_after}" / "unrender.sqlite3")
    database.initialize()
    with database.connect() as conn:
        conn.execute("DROP INDEX api_idempotency_expiry_idx")
        conn.execute("DROP TABLE api_idempotency")
        conn.execute(
            "CREATE TABLE api_idempotency ("
            "user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,"
            "idempotency_key TEXT NOT NULL,request_sha256 TEXT NOT NULL,response_json TEXT,"
            "created_at TEXT NOT NULL,completed_at TEXT,PRIMARY KEY(user_id,idempotency_key))"
        )
        conn.execute(
            "INSERT INTO users(id,email,password_hash,credit_balance,created_at) "
            "VALUES ('crash-user','crash@example.com','hash',0,'2026-01-01T00:00:00Z')"
        )
        conn.execute(
            "INSERT INTO api_idempotency VALUES "
            "('crash-user','key','sha','{}','2026-01-01T00:00:00Z',NULL)"
        )
        conn.execute("UPDATE schema_meta SET version=4")
    statements = 0

    def fault() -> None:
        nonlocal statements
        statements += 1
        if statements == crash_after:
            raise RuntimeError("injected statement crash")

    database._migration_fault_hook = fault
    with pytest.raises(RuntimeError, match="statement crash"):
        database.initialize()
    with database.connect() as conn:
        assert conn.execute("SELECT version FROM schema_meta").fetchone()[0] == 4
        assert (
            conn.execute(
                "SELECT response_json FROM api_idempotency WHERE user_id='crash-user'"
            ).fetchone()[0]
            == "{}"
        )
        assert (
            conn.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='api_idempotency_v5_new'"
            ).fetchone()
            is None
        )
        assert conn.execute("PRAGMA foreign_key_check").fetchone() is None
    database._migration_fault_hook = None
    database.initialize()
    with database.connect() as conn:
        row = conn.execute("SELECT * FROM api_idempotency").fetchone()
        assert conn.execute("SELECT version FROM schema_meta").fetchone()[0] == SCHEMA_VERSION
        assert row["response_json"] == "{}"
        assert row["expires_at"] == "2026-01-31T00:00:00.000Z"


@pytest.mark.parametrize("crash_after", range(1, 9))
def test_v6_to_v7_capacity_and_session_migration_is_atomic(
    tmp_path: Path, crash_after: int
) -> None:
    database = Database(tmp_path / f"migration-v7-{crash_after}" / "unrender.sqlite3")
    database.initialize()
    with database.connect() as conn:
        conn.execute(
            "INSERT INTO users(id,email,password_hash,credit_balance,created_at) "
            "VALUES ('v7-user','v7@example.com','hash',0,'2026-01-01T00:00:00Z')"
        )
        conn.execute("DROP TABLE storage_reservations")
        for table, column in (
            ("users", "session_generation"),
            ("sessions", "session_generation"),
            ("jobs", "result_reservation_bytes"),
            ("jobs", "result_reservation_attempt"),
            ("jobs", "retained_byte_reservation"),
        ):
            conn.execute(f'ALTER TABLE "{table}" DROP COLUMN "{column}"')
        conn.execute("UPDATE schema_meta SET version=6")
    statements = 0

    def fault() -> None:
        nonlocal statements
        statements += 1
        if statements == crash_after:
            raise RuntimeError("injected v7 statement crash")

    database._migration_fault_hook = fault
    with pytest.raises(RuntimeError, match="v7 statement crash"):
        database.initialize()
    with database.connect() as conn:
        assert conn.execute("SELECT version FROM schema_meta").fetchone()[0] == 6
        assert "session_generation" not in Database._columns(conn, "users")
        assert "result_reservation_bytes" not in Database._columns(conn, "jobs")
        assert (
            conn.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='storage_reservations'"
            ).fetchone()
            is None
        )
        assert conn.execute("SELECT email FROM users WHERE id='v7-user'").fetchone()[0] == (
            "v7@example.com"
        )
    database._migration_fault_hook = None
    database.initialize()
    database.initialize()
    with database.connect() as conn:
        assert conn.execute("SELECT version FROM schema_meta").fetchone()[0] == SCHEMA_VERSION
        assert "session_generation" in Database._columns(conn, "users")
        assert "result_reservation_bytes" in Database._columns(conn, "jobs")
        assert (
            conn.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='storage_reservations'"
            ).fetchone()
            is not None
        )
        assert conn.execute("PRAGMA foreign_key_check").fetchone() is None


@pytest.mark.parametrize("crash_after", range(1, 4))
def test_v7_to_v8_storage_owner_tokens_are_atomic_and_opaque(
    tmp_path: Path, crash_after: int
) -> None:
    database = Database(tmp_path / f"migration-v8-{crash_after}" / "unrender.sqlite3")
    database.initialize()
    with database.connect() as conn:
        conn.execute(
            "INSERT INTO users(id,email,password_hash,credit_balance,created_at) "
            "VALUES ('v8-user','v8@example.com','hash',0,'2026-01-01T00:00:00Z')"
        )
        conn.execute(
            "INSERT INTO storage_reservations("
            "id,owner_token,user_id,kind,byte_count,storage_path,created_at,expires_at"
            ") VALUES ('legacy-reservation','old-token','v8-user','staging',9,"
            "'/legacy/staging.tmp','2026-01-01T00:00:00Z','2099-01-01T00:00:00Z')"
        )
        conn.execute("ALTER TABLE storage_reservations DROP COLUMN owner_token")
        conn.execute("UPDATE schema_meta SET version=7")
    statements = 0

    def fault() -> None:
        nonlocal statements
        statements += 1
        if statements == crash_after:
            raise RuntimeError("injected v8 statement crash")

    database._migration_fault_hook = fault
    with pytest.raises(RuntimeError, match="v8 statement crash"):
        database.initialize()
    with database.connect() as conn:
        assert conn.execute("SELECT version FROM schema_meta").fetchone()[0] == 7
        assert "owner_token" not in Database._columns(conn, "storage_reservations")
        assert (
            conn.execute(
                "SELECT byte_count FROM storage_reservations WHERE id='legacy-reservation'"
            ).fetchone()[0]
            == 9
        )
    database._migration_fault_hook = None
    database.initialize()
    database.initialize()
    with database.connect() as conn:
        migrated = conn.execute(
            "SELECT owner_token,byte_count FROM storage_reservations WHERE id='legacy-reservation'"
        ).fetchone()
        assert conn.execute("SELECT version FROM schema_meta").fetchone()[0] == SCHEMA_VERSION
        assert migrated["byte_count"] == 9
        assert re.fullmatch(r"[0-9a-f]{48}", migrated["owner_token"])
        assert conn.execute("PRAGMA foreign_key_check").fetchone() is None


def test_v4_partial_rename_shape_is_resumed_without_losing_idempotency_rows(
    tmp_path: Path,
) -> None:
    database = Database(tmp_path / "migration-partial-v4" / "unrender.sqlite3")
    database.initialize()
    with database.connect() as conn:
        conn.execute(
            "INSERT INTO users(id,email,password_hash,credit_balance,created_at) "
            "VALUES ('partial-user','partial@example.com','hash',0,'2026-01-01T00:00:00Z')"
        )
        conn.execute(
            "INSERT INTO api_idempotency("
            "user_id,idempotency_key,request_sha256,response_json,created_at,completed_at,"
            "expires_at,expired_at) VALUES ("
            "'partial-user','key','sha','{}','2026-01-01T00:00:00Z',NULL,"
            "'2026-02-01T00:00:00Z',NULL)"
        )
        conn.execute("DROP INDEX api_idempotency_expiry_idx")
        conn.execute("ALTER TABLE api_idempotency RENAME TO api_idempotency_v4")
        conn.execute("UPDATE schema_meta SET version=4")
    database.initialize()
    database.initialize()
    with database.connect() as conn:
        row = conn.execute("SELECT * FROM api_idempotency").fetchone()
        assert conn.execute("SELECT version FROM schema_meta").fetchone()[0] == SCHEMA_VERSION
        assert row["user_id"] == "partial-user"
        assert row["expires_at"] == "2026-02-01T00:00:00Z"
        assert conn.execute("PRAGMA foreign_key_check").fetchone() is None


def test_backup_is_coordinated_hash_verified_and_rewrites_restored_paths(
    tmp_path: Path,
) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=1)
    user_id = customer_id(service, "backup@example.com")
    job = _paid_job(service, user_id, color="orange")
    uncommitted = service.settings.storage_dir / "jobs" / "orphan" / "not-in-database.source"
    uncommitted.parent.mkdir(parents=True)
    uncommitted.write_bytes(b"must not enter recovery set")
    backup = create_backup(service.settings, tmp_path / "recovery-set")
    assert not (backup / "storage" / "jobs" / "orphan" / uncommitted.name).exists()
    restored = restore_backup(backup, tmp_path / "restored-data")
    restored_database = Database(restored / "unrender.sqlite3")
    with restored_database.connect() as conn:
        restored_job = conn.execute(
            "SELECT source_path FROM jobs WHERE id=?", (job["id"],)
        ).fetchone()
        assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert conn.execute("PRAGMA foreign_key_check").fetchone() is None
    assert str(restored_job["source_path"]).startswith(str(restored / "storage"))
    assert Path(restored_job["source_path"]).is_file()

    with (
        service.database.operational_lock(),
        pytest.raises(BackupError, match="mutation is active"),
    ):
        create_backup(service.settings, tmp_path / "blocked-recovery-set")
    assert not (tmp_path / "blocked-recovery-set").exists()

    stored_source = next((backup / "storage" / "jobs").rglob("*.source"))
    stored_source.write_bytes(b"tampered")
    with pytest.raises(BackupError, match="hash/size"):
        restore_backup(backup, tmp_path / "tampered-restore")

    backup = create_backup(service.settings, tmp_path / "path-tamper-recovery-set")
    manifest_path = backup / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["source_data_dir"] = "/wrong/live/root"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    with pytest.raises(BackupError, match="outside product storage"):
        restore_backup(backup, tmp_path / "path-tampered-restore")


def test_cursor_inventories_key_cap_revoke_all_and_audit_rollups(
    tmp_path: Path,
) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        initial_credits=8,
        max_active_api_keys_per_user=3,
        max_api_key_records_per_user=5,
        max_audit_events_per_job=5,
        max_audit_events_per_user=30,
    )
    user_id = customer_id(service, "inventory@example.com")
    upload = service.prepare_upload(
        user_id=user_id, filename="inventory.png", content=png_bytes(color="yellow")
    )
    jobs = [
        service.create_job(
            user_id=user_id,
            upload_id=str(upload["id"]),
            page_index=0,
            crop=None,
        )
        for _ in range(5)
    ]
    keys = [service.create_api_key(user_id=user_id, name=f"key {index}") for index in range(3)]
    with pytest.raises(ProductError, match="Revoke an active"):
        service.create_api_key(user_id=user_id, name="hidden fourth key")

    job_ids: list[str] = []
    cursor = None
    while True:
        page = service.list_jobs_page(user_id, cursor=cursor, limit=2)
        job_ids.extend(str(item["id"]) for item in page["items"])
        cursor = page["next_cursor"]
        if cursor is None:
            break
    assert set(job_ids) == {str(job["id"]) for job in jobs}

    prefixes: list[str] = []
    cursor = None
    while True:
        page = service.list_api_keys_page(user_id=user_id, cursor=cursor, limit=2)
        prefixes.extend(str(item["prefix"]) for item in page["items"])
        cursor = page["next_cursor"]
        if cursor is None:
            break
    assert set(prefixes) == {str(key["prefix"]) for key in keys}
    assert service.revoke_all_api_keys(user_id=user_id) == 3
    assert all(item["revoked_at"] for item in service.list_api_keys(user_id=user_id))

    audited_job = str(jobs[0]["id"])
    with service.database.transaction(immediate=True) as conn:
        for index in range(7):
            service._audit(
                conn,
                user_id=user_id,
                job_id=audited_job,
                event_type="adversarial_event",
                details={"index": index},
            )
    page = service.job_audit_page(user_id=user_id, job_id=audited_job, limit=2)
    raw_events = list(page["items"])
    archived = sum(int(item["count"]) for item in page["rollups"])
    cursor = page["next_cursor"]
    while cursor:
        page = service.job_audit_page(
            user_id=user_id,
            job_id=audited_job,
            cursor=cursor,
            limit=2,
        )
        raw_events.extend(page["items"])
        cursor = page["next_cursor"]
    assert len(raw_events) <= 5
    assert len(raw_events) + archived == 8  # job_queued plus seven injected events


def test_api_key_inventory_does_not_hide_the_101st_record(tmp_path: Path) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        max_active_api_keys_per_user=101,
        max_api_key_records_per_user=110,
        max_audit_events_per_user=500,
    )
    user_id = customer_id(service, "key-pagination@example.com")
    expected = {
        service.create_api_key(user_id=user_id, name=f"integration {index}")["prefix"]
        for index in range(101)
    }
    observed: set[str] = set()
    cursor = None
    while True:
        page = service.list_api_keys_page(user_id=user_id, cursor=cursor, limit=17)
        observed.update(str(item["prefix"]) for item in page["items"])
        cursor = page["next_cursor"]
        if cursor is None:
            break
    assert observed == expected
    assert service.revoke_all_api_keys(user_id=user_id) == 101


def test_sessions_and_account_audit_state_remain_bounded(tmp_path: Path) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=True,
        max_sessions_per_user=2,
        max_audit_events_per_job=3,
        max_audit_events_per_user=5,
    )
    user_id = customer_id(service, "bounded@example.com")
    for _ in range(5):
        service.create_session(user_id)
    with service.database.transaction(immediate=True) as conn:
        for _ in range(10):
            service._audit(conn, user_id=user_id, event_type="account_probe")
    with service.database.connect() as conn:
        assert (
            conn.execute("SELECT COUNT(*) FROM sessions WHERE user_id=?", (user_id,)).fetchone()[0]
            == 2
        )
        assert (
            conn.execute(
                "SELECT COUNT(*) FROM audit_events WHERE user_id=?", (user_id,)
            ).fetchone()[0]
            <= 5
        )
        rollup = conn.execute(
            "SELECT event_count FROM audit_rollups "
            "WHERE user_id=? AND job_scope='account' AND event_type='account_probe'",
            (user_id,),
        ).fetchone()
    assert rollup is not None and int(rollup["event_count"]) > 0

    demo = service.demo_session()
    demo_user = service.session_user(demo["session"])
    assert demo_user is not None
    demo_id = str(demo_user["id"])
    service.logout(demo["session"])
    with service.database.connect() as conn:
        assert conn.execute("SELECT 1 FROM users WHERE id=?", (demo_id,)).fetchone() is None
        assert (
            conn.execute("SELECT 1 FROM audit_events WHERE user_id=?", (demo_id,)).fetchone()
            is None
        )


def test_auth_kdf_concurrency_docs_errors_live_webhooks_and_lazy_import(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    app = create_app(
        settings_for(
            tmp_path,
            seed_demo_account=False,
            auth_rate_limit_per_minute=100,
            max_concurrent_auth_requests=1,
            stripe_secret_key="sk_test_local",
            stripe_webhook_secret="whsec_local",
            stripe_price_id="price_local",
        )
    )
    entered = threading.Event()
    release = threading.Event()

    def slow_authenticate(_: str, __: str) -> dict[str, str]:
        entered.set()
        assert release.wait(timeout=5)
        raise ProductError("invalid_credentials", "Email or password is incorrect", 401)

    app.state.service.authenticate = slow_authenticate

    @app.get("/unexpected-test-error")
    def unexpected_test_error() -> None:
        raise RuntimeError("private detail")

    with TestClient(app, raise_server_exceptions=False) as client:
        first_status: list[int] = []

        def first_login() -> None:
            response = client.post(
                "/api/auth/login",
                json={"email": "first@example.com", "password": "password"},
            )
            first_status.append(response.status_code)

        thread = threading.Thread(target=first_login)
        thread.start()
        assert entered.wait(timeout=5)
        blocked = client.post(
            "/api/auth/login",
            json={"email": "second@example.com", "password": "password"},
        )
        assert blocked.status_code == 503
        assert blocked.json()["error"]["code"] == "auth_capacity_reached"
        release.set()
        thread.join(timeout=5)
        assert first_status == [401]
        assert client.get("/docs").status_code == 404
        assert client.get("/openapi.json").status_code == 404
        failed = client.get("/unexpected-test-error")
        assert failed.status_code == 500
        assert failed.json() == {
            "error": {"code": "internal_error", "message": "The request could not be completed"}
        }
        assert failed.headers["X-Content-Type-Options"] == "nosniff"
        assert "default-src 'self'" in failed.headers["Content-Security-Policy"]

        live_event = {
            "id": "evt_live_forbidden",
            "livemode": True,
            "type": "checkout.session.completed",
            "data": {"object": {"livemode": True}},
        }
        monkeypatch.setattr(stripe.Webhook, "construct_event", lambda *_: live_event)
        live = client.post(
            "/api/billing/webhook",
            content=b"live-payload",
            headers={"Stripe-Signature": "test-signature"},
        )
        assert live.status_code == 422
        assert live.json()["error"]["code"] == "billing_live_event_rejected"

    probe = subprocess.run(
        [
            sys.executable,
            "-c",
            "import pathlib,sys; import unrender.product; "
            "assert 'unrender.product.web' not in sys.modules; "
            "assert 'unrender.product.storage' not in sys.modules; "
            "assert not pathlib.Path('.unrender-data').exists()",
        ],
        cwd=tmp_path,
        env={**os.environ, "PYTHONPATH": str(STATIC_DIR.parents[2])},
        capture_output=True,
        text=True,
        check=False,
    )
    assert probe.returncode == 0, probe.stderr


def test_outer_asgi_failure_is_sanitized_with_security_headers() -> None:
    async def failing_app(_scope: object, _receive: object, _send: object) -> None:
        raise RuntimeError("private middleware detail")

    messages: list[dict[str, object]] = []

    async def receive() -> dict[str, object]:
        return {"type": "http.request", "body": b"", "more_body": False}

    async def send(message: dict[str, object]) -> None:
        messages.append(message)

    middleware = SecurityHeadersMiddleware(failing_app, secure_cookies=True)  # type: ignore[arg-type]
    asyncio.run(middleware({"type": "http", "path": "/failure"}, receive, send))  # type: ignore[arg-type]
    start = next(message for message in messages if message["type"] == "http.response.start")
    headers = {
        bytes(name).decode("latin-1").casefold(): bytes(value).decode("latin-1")
        for name, value in start["headers"]  # type: ignore[union-attr]
    }
    body = b"".join(
        bytes(message.get("body", b""))
        for message in messages
        if message["type"] == "http.response.body"
    )
    assert start["status"] == 500
    assert headers["x-content-type-options"] == "nosniff"
    assert "default-src 'self'" in headers["content-security-policy"]
    assert b"private middleware detail" not in body


def test_expensive_upload_concurrency_is_rejected_before_route_processing(
    tmp_path: Path,
) -> None:
    app = create_app(
        settings_for(
            tmp_path,
            seed_demo_account=False,
            max_concurrent_expensive_requests=1,
        )
    )
    entered = threading.Event()
    release = threading.Event()
    route_calls = 0

    def slow_upload(**_: object) -> dict[str, object]:
        nonlocal route_calls
        route_calls += 1
        entered.set()
        assert release.wait(timeout=5)
        return {
            "id": "bounded-upload",
            "name": "chart.png",
            "mime_type": "image/png",
            "page_count": 1,
            "expires_at": timestamp(),
        }

    app.state.service.prepare_upload_stream = slow_upload
    with TestClient(app) as client:
        registered = client.post(
            "/api/auth/register",
            json={"email": "upload-limit@example.com", "password": "long enough password"},
        )
        assert registered.status_code == 201
        first_status: list[int] = []

        def first_upload() -> None:
            response = client.post(
                "/api/uploads",
                headers=csrf_headers(client),
                files={"file": ("first.png", png_bytes(), "image/png")},
            )
            first_status.append(response.status_code)

        thread = threading.Thread(target=first_upload)
        thread.start()
        assert entered.wait(timeout=5)
        blocked = client.post(
            "/api/uploads",
            headers=csrf_headers(client),
            files={"file": ("second.png", png_bytes(), "image/png")},
        )
        assert blocked.status_code == 503
        assert blocked.json()["error"]["code"] == "expensive_capacity_reached"
        release.set()
        thread.join(timeout=5)
    assert first_status == [201]
    assert route_calls == 1


def test_modal_contract_is_exact_nonspending_and_release_digest_is_case_normalized(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    resolved: list[tuple[str, str]] = []

    hydrated = []

    async def hydrate():
        hydrated.append(True)

    async def forbidden(*args):
        raise AssertionError("spent")

    def from_name(app_name: str, function_name: str) -> SimpleNamespace:
        resolved.append((app_name, function_name))
        return SimpleNamespace(
            remote=SimpleNamespace(aio=forbidden), hydrate=SimpleNamespace(aio=hydrate)
        )

    monkeypatch.setitem(
        sys.modules,
        "modal",
        SimpleNamespace(Function=SimpleNamespace(from_name=from_name)),
    )
    settings = settings_for(
        tmp_path,
        extractor_backend="modal",
        modal_provider_release="A" * 64,
    )
    extractor = ModalExtractor(settings)
    assert extractor.canary_contract()
    assert resolved == [("unrender-production", "infer_one")]
    assert hydrated == [True]

    response = {
        "json": json.loads(
            (STATIC_DIR / "demo" / "budget-quarter-result.json").read_text(encoding="utf-8")
        )["result"],
        "raw": "case-normalized",
        "provider_release": "a" * 64,
    }

    async def remote_response(*args):
        return response

    monkeypatch.setitem(
        sys.modules,
        "modal",
        SimpleNamespace(
            Function=SimpleNamespace(
                from_name=lambda *_: SimpleNamespace(remote=SimpleNamespace(aio=remote_response))
            )
        ),
    )
    assert extractor.extract(png_bytes()).model_version.endswith("+provider:aaaaaaaaaaaa")


def test_nonfinite_chart_coordinates_are_rejected() -> None:
    payload = {
        "chart_type": "line",
        "title": "Nonfinite",
        "x_axis": {"label": "x", "unit": None},
        "y_axis": {"label": "y", "unit": None},
        "series": [{"name": None, "points": [{"x": float("nan"), "y": 1.0}]}],
    }
    with pytest.raises(ValueError):
        ChartData.model_validate(payload)
    payload["series"][0]["points"][0] = {"x": 1.0, "y": float("inf")}
    with pytest.raises(ValueError):
        ChartData.model_validate(payload)


def test_correction_transport_accepts_max_contract_and_rejects_exact_overflow(
    tmp_path: Path,
) -> None:
    settings = settings_for(tmp_path, seed_demo_account=False)
    app = create_app(settings)
    with TestClient(app) as client:
        registered = client.post(
            "/api/auth/register",
            json={"email": "large-result@example.com", "password": "long result password"},
        )
        assert registered.status_code == 201
        upload = client.post(
            "/api/uploads",
            headers=csrf_headers(client),
            files={"file": ("large.png", png_bytes(color="navy"), "image/png")},
        )
        job = client.post(
            "/api/jobs",
            headers={**csrf_headers(client), "Idempotency-Key": "large-result-job"},
            json={"upload_id": upload.json()["id"], "page_index": 0},
        )
        assert job.status_code == 202
        _install_successful_extractor(app.state.service)
        assert app.state.service.process_one()
        job_id = job.json()["id"]
        result = {
            "chart_type": "line",
            "title": "Near transport maximum",
            "x_axis": {"label": "x", "unit": None},
            "y_axis": {"label": "y", "unit": None},
            "series": [
                {
                    "name": "series",
                    "points": [
                        {"x": f"{index:05d}{'x' * 75}", "y": index} for index in range(10_000)
                    ],
                }
            ],
        }
        body = json.dumps({"result": result}, separators=(",", ":")).encode()
        assert 970_000 < len(body) <= settings.result_request_bytes
        accepted = client.patch(
            f"/api/jobs/{job_id}/result",
            headers={**csrf_headers(client), "Content-Type": "application/json"},
            content=body,
        )
        assert accepted.status_code == 200

        exact = b"{}" + (b" " * (settings.result_request_bytes - 2))
        assert len(exact) == settings.result_request_bytes
        exact_response = client.patch(
            f"/api/jobs/{job_id}/result",
            headers={**csrf_headers(client), "Content-Type": "application/json"},
            content=exact,
        )
        assert exact_response.status_code == 422
        overflow = exact + b" "
        rejected = client.patch(
            f"/api/jobs/{job_id}/result",
            headers={**csrf_headers(client), "Content-Type": "application/json"},
            content=overflow,
        )
        assert rejected.status_code == 413
        assert rejected.json()["error"]["code"] == "request_too_large"


def test_browser_job_idempotency_replays_concurrently_without_double_charge(
    tmp_path: Path,
) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=3)
    user_id = customer_id(service, "browser-idempotency@example.com")
    upload = service.prepare_upload(
        user_id=user_id, filename="idempotent.png", content=png_bytes(color="green")
    )
    barrier = threading.Barrier(3)
    results: list[dict[str, Any]] = []
    errors: list[Exception] = []

    def submit() -> None:
        barrier.wait()
        try:
            results.append(
                service.create_job(
                    user_id=user_id,
                    upload_id=str(upload["id"]),
                    page_index=0,
                    crop=None,
                    idempotency_key="browser-concurrent-0001",
                )
            )
        except Exception as exc:  # pragma: no cover - assertion exposes exact race failure
            errors.append(exc)

    threads = [threading.Thread(target=submit) for _ in range(2)]
    for thread in threads:
        thread.start()
    barrier.wait()
    for thread in threads:
        thread.join(timeout=10)
    assert all(not thread.is_alive() for thread in threads)
    assert errors == []
    assert len(results) == 2 and results[0] == results[1]
    assert service.account(user_id)["credits"] == 2
    with pytest.raises(ProductError) as conflict:
        service.create_job(
            user_id=user_id,
            upload_id=str(upload["id"]),
            page_index=0,
            crop={"x": 0.1, "y": 0.1, "width": 0.5, "height": 0.5},
            idempotency_key="browser-concurrent-0001",
        )
    assert conflict.value.code == "idempotency_conflict"
    with service.database.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM jobs").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM api_idempotency").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM storage_reservations").fetchone()[0] == 0
    with service.database.transaction(immediate=True) as conn:
        conn.execute("DELETE FROM uploads WHERE id=?", (upload["id"],))
    assert (
        service.create_job(
            user_id=user_id,
            upload_id=str(upload["id"]),
            page_index=0,
            crop=None,
            idempotency_key="browser-concurrent-0001",
        )
        == results[0]
    )


def test_ledger_reserves_every_outstanding_refund_through_all_terminal_paths(
    tmp_path: Path,
) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        initial_credits=5,
        max_credit_ledger_records_per_user=9,
        max_recovery_attempts=0,
    )
    user_id = customer_id(service, "ledger-reserve@example.com")
    jobs = [_paid_job(service, user_id, color=color) for color in ("red", "blue", "green", "black")]
    with pytest.raises(ProductError, match="Credit activity is paused"):
        _paid_job(service, user_id, color="white")

    service.cancel(user_id=user_id, job_id=str(jobs[0]["id"]))
    failed_claim = service.claim_next_job("ledger-failure")
    assert failed_claim is not None
    assert service._finish_failed_claim(failed_claim, "predispatch_failure", "safe refund")
    expired_claim = service.claim_next_job("ledger-recovery")
    assert expired_claim is not None
    with service.database.transaction(immediate=True) as conn:
        conn.execute(
            "UPDATE jobs SET lease_expires_at=? WHERE id=?",
            (timestamp(utcnow() - timedelta(seconds=1)), expired_claim.job_id),
        )
    assert service.recover_interrupted_jobs() == 1
    service.cancel(user_id=user_id, job_id=str(jobs[3]["id"]))

    assert service.account(user_id)["credits"] == 5
    with service.database.connect() as conn:
        ledger = conn.execute(
            "SELECT reason FROM credit_ledger WHERE user_id=? ORDER BY created_at,id", (user_id,)
        ).fetchall()
        assert len(ledger) == 9
        assert sum("refund" in row["reason"] for row in ledger) == 4
    with pytest.raises(ProductError, match="Credit activity is paused"):
        service.apply_billing_event(
            event_id="evt_ledger_full",
            event_type="checkout.session.completed",
            user_id=user_id,
            credits=service.settings.credit_pack_size,
            payload_sha256="f" * 64,
        )
    with service.database.connect() as conn:
        assert (
            conn.execute("SELECT 1 FROM billing_events WHERE event_id='evt_ledger_full'").fetchone()
            is None
        )


def test_result_capacity_is_attempt_reserved_settled_and_released_predispatch(
    tmp_path: Path,
) -> None:
    service = service_for(
        tmp_path / "queued",
        seed_demo_account=False,
        initial_credits=3,
        max_history_bytes_per_user=2_000_000,
    )
    user_id = customer_id(service, "result-reserve@example.com")
    upload = service.prepare_upload(
        user_id=user_id, filename="reserve.png", content=png_bytes(color="purple")
    )
    first = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    second = service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    with pytest.raises(ProductError, match="maximum-size result"):
        service.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    assert service.account(user_id)["credits"] == 1
    with service.database.connect() as conn:
        reservations = conn.execute(
            "SELECT attempt,result_reservation_bytes,result_reservation_attempt,"
            "retained_byte_reservation FROM jobs ORDER BY created_at"
        ).fetchall()
    assert len(reservations) == 2
    assert all(row["result_reservation_attempt"] == row["attempt"] == 1 for row in reservations)
    assert all(row["result_reservation_bytes"] == 1_000_000 for row in reservations)
    service.cancel(user_id=user_id, job_id=str(first["id"]))
    replacement = service.create_job(
        user_id=user_id, upload_id=upload["id"], page_index=0, crop=None
    )
    assert replacement["status"] == "queued"
    service.cancel(user_id=user_id, job_id=str(second["id"]))
    service.cancel(user_id=user_id, job_id=str(replacement["id"]))

    reprocess_service = service_for(
        tmp_path / "reprocess",
        seed_demo_account=False,
        initial_credits=3,
        max_history_bytes_per_user=1_000_000,
    )
    reprocess_user = customer_id(reprocess_service, "reprocess-reserve@example.com")
    completed = _paid_job(reprocess_service, reprocess_user, color="orange")
    _install_successful_extractor(reprocess_service)
    successful_extract = reprocess_service.extractor.extract
    reprocess_service.extractor = SimpleNamespace(
        extract=lambda image: ExtractionOutput(
            chart=successful_extract(image).chart,
            raw="é" * 600_000,
            extractor="test",
            model_version="test-pinned",
        )
    )
    assert reprocess_service.process_one()
    assert reprocess_service.account(reprocess_user)["credits"] == 2
    with reprocess_service.database.connect() as conn:
        settled = conn.execute(
            "SELECT result_reservation_bytes,result_reservation_attempt,"
            "retained_byte_reservation,attempt,LENGTH(CAST(raw_result AS BLOB)) AS raw_bytes "
            "FROM jobs WHERE id=?",
            (completed["id"],),
        ).fetchone()
    assert tuple(settled) == (0, None, 0, 1, reprocess_service.settings.max_result_json_bytes)
    with pytest.raises(ProductError, match="maximum-size result"):
        reprocess_service.reprocess(user_id=reprocess_user, job_id=str(completed["id"]))
    with pytest.raises(ProductError, match="maximum-size result"):
        reprocess_service.submit_api_extraction(
            user_id=reprocess_user,
            filename="second.png",
            content=png_bytes(color="blue"),
            page_index=0,
            idempotency_key="result-capacity-api",
        )
    assert reprocess_service.account(reprocess_user)["credits"] == 2
    with reprocess_service.database.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM jobs").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM storage_reservations").fetchone()[0] == 0


def test_crash_durable_publication_reservations_reconcile_staging_and_namespaces(
    tmp_path: Path,
) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        initial_credits=2,
        reconciliation_grace_seconds=0,
    )
    user_id = customer_id(service, "publication-crash@example.com")

    def crash_after_staging(stage: str) -> None:
        if stage == "staging_file_durable":
            raise SystemExit("simulated staging publication loss")

    service.storage._publication_fault_hook = crash_after_staging
    with pytest.raises(SystemExit, match="staging publication loss"):
        service.prepare_upload(
            user_id=user_id, filename="staging-crash.png", content=png_bytes(color="black")
        )
    service.storage._publication_fault_hook = None
    with service.database.connect() as conn:
        staging_reservation = conn.execute(
            "SELECT id,storage_path FROM storage_reservations WHERE kind='staging'"
        ).fetchone()
    staged_path = Path(staging_reservation["storage_path"])
    assert staged_path.is_file()

    peer = ProductService(
        settings=service.settings,
        database=Database(service.settings.database_path),
        storage=Storage(service.settings),
        extractor=ReplayExtractor(STATIC_DIR),
        static_dir=STATIC_DIR,
    )
    peer.initialize()
    assert staged_path.is_file()
    with peer.database.transaction(immediate=True) as conn:
        conn.execute(
            "UPDATE storage_reservations SET expires_at=? WHERE id=?",
            (timestamp(utcnow() - timedelta(seconds=1)), staging_reservation["id"]),
        )
    assert peer.cleanup_expired()["storage_reservations"] == 1
    assert not staged_path.exists()

    def crash_after_upload(stage: str) -> None:
        if stage == "upload_published":
            raise SystemExit("simulated process loss")

    service.storage._publication_fault_hook = crash_after_upload
    with pytest.raises(SystemExit, match="process loss"):
        service.prepare_upload(
            user_id=user_id, filename="crash.png", content=png_bytes(color="red")
        )
    service.storage._publication_fault_hook = None
    with service.database.connect() as conn:
        reservation = conn.execute(
            "SELECT id,storage_path FROM storage_reservations WHERE kind='upload'"
        ).fetchone()
        assert conn.execute("SELECT COUNT(*) FROM uploads").fetchone()[0] == 0
    published = Path(reservation["storage_path"])
    assert published.is_file()

    peer.initialize()
    assert published.is_file()
    with peer.database.transaction(immediate=True) as conn:
        conn.execute(
            "UPDATE storage_reservations SET expires_at=? WHERE id=?",
            (timestamp(utcnow() - timedelta(seconds=1)), reservation["id"]),
        )
    assert peer.cleanup_expired()["storage_reservations"] == 1
    assert not published.exists()

    upload = peer.prepare_upload(
        user_id=user_id, filename="job-crash.png", content=png_bytes(color="green")
    )

    def crash_after_job(stage: str) -> None:
        if stage == "job_copy_published":
            raise SystemExit("simulated job publication loss")

    peer.storage._publication_fault_hook = crash_after_job
    with pytest.raises(SystemExit, match="job publication loss"):
        peer.create_job(user_id=user_id, upload_id=upload["id"], page_index=0, crop=None)
    peer.storage._publication_fault_hook = None
    with peer.database.connect() as conn:
        job_reservation = conn.execute(
            "SELECT id,storage_path FROM storage_reservations WHERE kind='job_copy'"
        ).fetchone()
        assert conn.execute("SELECT COUNT(*) FROM jobs").fetchone()[0] == 0
    job_copy = Path(job_reservation["storage_path"])
    assert job_copy.is_file()
    with peer.database.transaction(immediate=True) as conn:
        conn.execute(
            "UPDATE storage_reservations SET expires_at=? WHERE id=?",
            (timestamp(utcnow() - timedelta(seconds=1)), job_reservation["id"]),
        )
    assert peer.cleanup_expired()["storage_reservations"] == 1
    assert not job_copy.exists()
    assert peer.upload_preview(user_id=user_id, upload_id=upload["id"], page_index=0)


def test_storage_reservation_tokens_renew_and_expired_publication_fails_closed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=1)
    user_id = customer_id(service, "reservation-owner@example.com")
    reservation_path = service.storage.root / "staging" / "owned.tmp"
    reservation = service._reserve_storage_bytes(
        user_id=user_id,
        byte_count=1,
        kind="staging",
        storage_path=reservation_path,
    )
    with service.database.transaction(immediate=True) as conn:
        short_expiry = timestamp(utcnow() + timedelta(seconds=1))
        conn.execute(
            "UPDATE storage_reservations SET expires_at=? WHERE id=?",
            (short_expiry, reservation.reservation_id),
        )

    stale = StorageReservation(reservation.reservation_id, "stale-owner-token")
    with pytest.raises(ProductError) as stale_resize:
        service._resize_storage_reservation(
            user_id=user_id,
            reservation=stale,
            byte_count=2,
            kind="staging",
            storage_path=reservation_path,
        )
    assert stale_resize.value.code == "storage_reservation_lost"
    service._release_storage_reservation(user_id=user_id, reservation=stale)
    service._resize_storage_reservation(
        user_id=user_id,
        reservation=reservation,
        byte_count=2,
        kind="staging",
        storage_path=reservation_path,
    )
    with service.database.connect() as conn:
        renewed = conn.execute(
            "SELECT owner_token,byte_count,expires_at FROM storage_reservations WHERE id=?",
            (reservation.reservation_id,),
        ).fetchone()
    assert renewed["owner_token"] == reservation.owner_token
    assert renewed["byte_count"] == 2
    assert renewed["expires_at"] > short_expiry
    service._release_storage_reservation(user_id=user_id, reservation=reservation)

    peer = ProductService(
        settings=service.settings,
        database=Database(service.settings.database_path),
        storage=Storage(service.settings),
        extractor=ReplayExtractor(STATIC_DIR),
        static_dir=STATIC_DIR,
    )
    peer.initialize()
    original_resize = service._resize_storage_reservation

    def expire_before_publication(**kwargs: Any) -> None:
        original_resize(**kwargs)
        if kwargs["kind"] != "upload":
            return
        live = kwargs["reservation"]
        with peer.database.transaction(immediate=True) as conn:
            conn.execute(
                "UPDATE storage_reservations SET expires_at=? WHERE id=? AND owner_token=?",
                (
                    timestamp(utcnow() - timedelta(seconds=1)),
                    live.reservation_id,
                    live.owner_token,
                ),
            )
        assert peer.cleanup_storage_reservations() == 1

    monkeypatch.setattr(service, "_resize_storage_reservation", expire_before_publication)
    with pytest.raises(ProductError) as expired:
        service.prepare_upload(
            user_id=user_id,
            filename="expired-publication.png",
            content=png_bytes(color="red"),
        )
    assert expired.value.code == "storage_reservation_lost"
    with service.database.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM uploads").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM storage_reservations").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM pending_deletions").fetchone()[0] == 0
    assert service.storage.object_paths() == []


def test_staging_token_loss_removes_written_file_before_returning_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=1)
    user_id = customer_id(service, "staging-token-loss@example.com")
    peer = ProductService(
        settings=service.settings,
        database=Database(service.settings.database_path),
        storage=Storage(service.settings),
        extractor=ReplayExtractor(STATIC_DIR),
        static_dir=STATIC_DIR,
    )
    peer.initialize()
    original_stage = service.storage.stage_upload
    written: list[Path] = []

    def lose_reservation_after_write(source: Any, **kwargs: Any):
        staged = original_stage(source, **kwargs)
        written.append(staged.path)
        with peer.database.transaction(immediate=True) as conn:
            assert (
                conn.execute(
                    "DELETE FROM storage_reservations WHERE storage_path=?",
                    (str(staged.path),),
                ).rowcount
                == 1
            )
        return staged

    monkeypatch.setattr(service.storage, "stage_upload", lose_reservation_after_write)
    with pytest.raises(ProductError) as lost:
        service.prepare_upload(
            user_id=user_id,
            filename="lost-token.png",
            content=png_bytes(color="orange"),
        )
    assert lost.value.code == "storage_reservation_lost"
    assert written and all(not path.exists() for path in written)
    with service.database.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM storage_reservations").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM pending_deletions").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM uploads").fetchone()[0] == 0


def test_publication_and_deletion_reference_check_share_one_cross_process_fence(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=1)
    user_id = customer_id(service, "publication-delete-race@example.com")
    peer = ProductService(
        settings=service.settings,
        database=Database(service.settings.database_path),
        storage=Storage(service.settings),
        extractor=ReplayExtractor(STATIC_DIR),
        static_dir=STATIC_DIR,
    )
    peer.initialize()

    expected_path: list[Path] = []
    original_resize = service._resize_storage_reservation

    def queue_stale_deletion(**kwargs: Any) -> None:
        original_resize(**kwargs)
        if kwargs["kind"] != "upload":
            return
        path = Path(kwargs["storage_path"])
        expected_path.append(path)
        with peer.database.transaction(immediate=True) as conn:
            peer._queue_deletion(
                conn,
                path,
                "deterministic_publication_race",
                user_id=user_id,
                byte_size=int(kwargs["byte_count"]),
            )

    monkeypatch.setattr(service, "_resize_storage_reservation", queue_stale_deletion)
    exclusive_attempted = threading.Event()
    original_lock = peer.database.operational_lock

    def tracked_lock(*, exclusive: bool = False, timeout_seconds: float = 10.0):
        if exclusive:
            exclusive_attempted.set()
        return original_lock(exclusive=exclusive, timeout_seconds=timeout_seconds)

    monkeypatch.setattr(peer.database, "operational_lock", tracked_lock)
    drain_results: list[int] = []
    drain_threads: list[threading.Thread] = []

    def publication_hook(stage: str) -> None:
        if stage != "upload_published":
            return
        thread = threading.Thread(
            target=lambda: drain_results.append(
                peer.drain_deletion_queue(paths=[str(expected_path[0])])
            )
        )
        drain_threads.append(thread)
        thread.start()
        assert exclusive_attempted.wait(timeout=5)
        time.sleep(0.05)
        assert thread.is_alive(), "exclusive deletion must wait for publication's shared fence"

    service.storage._publication_fault_hook = publication_hook
    uploaded = service.prepare_upload(
        user_id=user_id,
        filename="committed-race.png",
        content=png_bytes(color="green"),
    )
    service.storage._publication_fault_hook = None
    for thread in drain_threads:
        thread.join(timeout=5)
        assert not thread.is_alive()
    assert drain_results == [0]
    row = service.upload(user_id=user_id, upload_id=str(uploaded["id"]))
    assert Path(row["storage_path"]) == expected_path[0]
    assert expected_path[0].is_file()
    with service.database.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM pending_deletions").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM storage_reservations").fetchone()[0] == 0


def test_startup_never_invents_missing_result_capacity_before_provider_dispatch(
    tmp_path: Path,
) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=1)
    user_id = customer_id(service, "legacy-capacity@example.com")
    job = _paid_job(service, user_id, color="purple")
    with service.database.transaction(immediate=True) as conn:
        conn.execute(
            "UPDATE jobs SET result_reservation_bytes=0,result_reservation_attempt=NULL,"
            "retained_byte_reservation=0 WHERE id=?",
            (job["id"],),
        )

    class MustNotDispatch:
        def extract(self, _: bytes) -> None:
            raise AssertionError("provider dispatch must remain behind durable capacity")

    peer = ProductService(
        settings=service.settings,
        database=Database(service.settings.database_path),
        storage=Storage(service.settings),
        extractor=MustNotDispatch(),
        static_dir=STATIC_DIR,
    )
    peer.initialize(recover_jobs=False)
    assert peer.process_one("legacy-capacity-worker")
    failed = peer.get_job(user_id=user_id, job_id=str(job["id"]))
    assert failed["status"] == "failed"
    assert failed["error"]["code"] == "result_capacity_reservation_lost"
    assert peer.account(user_id)["credits"] == 1
    assert not peer.process_one("legacy-capacity-worker-retry")
    assert peer.account(user_id)["credits"] == 1
    with peer.database.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM provider_attempts").fetchone()[0] == 0
        assert (
            conn.execute(
                "SELECT COUNT(*) FROM credit_ledger WHERE idempotency_key=?",
                (f"job:{job['id']}:refund:1",),
            ).fetchone()[0]
            == 1
        )


def test_global_storage_reservations_are_cross_process_and_low_free_fail_closed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        initial_credits=2,
        max_upload_bytes=256 * 1024,
        min_free_storage_bytes=0,
    )
    first_user = customer_id(service, "storage-one@example.com")
    second_user = customer_id(service, "storage-two@example.com")
    with service.database.connect() as conn:
        baseline = service._retained_storage_bytes(conn)
    constrained = Settings(
        **{
            **service.settings.__dict__,
            "max_storage_bytes_global": baseline + service.settings.max_upload_bytes + 1024,
        }
    )
    service.settings = constrained
    service.storage.settings = constrained
    peer = ProductService(
        settings=constrained,
        database=Database(constrained.database_path),
        storage=Storage(constrained),
        extractor=ReplayExtractor(STATIC_DIR),
        static_dir=STATIC_DIR,
    )
    entered = threading.Event()
    release = threading.Event()
    original_stage = service.storage.stage_upload

    def slow_stage(source: Any, **kwargs: Any):
        entered.set()
        assert release.wait(timeout=5)
        return original_stage(source, **kwargs)

    monkeypatch.setattr(service.storage, "stage_upload", slow_stage)
    first_result: list[dict[str, Any]] = []

    def first_upload() -> None:
        first_result.append(
            service.prepare_upload(
                user_id=first_user, filename="one.png", content=png_bytes(color="red")
            )
        )

    thread = threading.Thread(target=first_upload)
    thread.start()
    assert entered.wait(timeout=5)
    with pytest.raises(ProductError) as capacity:
        peer.prepare_upload(
            user_id=second_user, filename="two.png", content=png_bytes(color="blue")
        )
    assert capacity.value.code == "service_storage_capacity_reached"
    release.set()
    thread.join(timeout=5)
    assert first_result and not thread.is_alive()

    relaxed = Settings(
        **{
            **constrained.__dict__,
            "max_storage_bytes_global": constrained.max_storage_bytes_global
            + constrained.max_upload_bytes,
        }
    )
    peer.settings = relaxed
    peer.storage.settings = relaxed

    monkeypatch.setattr(
        service_module.shutil,
        "disk_usage",
        lambda _: SimpleNamespace(
            free=constrained.min_free_storage_bytes + constrained.database_headroom_bytes - 1
        ),
    )
    with pytest.raises(ProductError) as low_free:
        peer.prepare_upload(
            user_id=second_user, filename="low-free.png", content=png_bytes(color="black")
        )
    assert low_free.value.code == "storage_free_space_guard"
    with service.database.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM storage_reservations").fetchone()[0] == 0


def test_database_row_admission_is_central_and_preserves_terminal_capacity(
    tmp_path: Path,
) -> None:
    source = (STATIC_DIR.parent / "service.py").read_text(encoding="utf-8")
    assert re.search(r"conn\.execute\(\s*\"INSERT INTO", source) is None

    service = service_for(tmp_path, seed_demo_account=False, initial_credits=1)
    user_id = customer_id(service, "row-pressure@example.com")
    job = _paid_job(service, user_id, color="yellow")
    with service.database.connect() as conn:
        current = service._database_row_count(conn, user_id)
        future = service._future_terminal_rows(conn, user_id)
    service.settings = Settings(
        **{
            **service.settings.__dict__,
            "max_database_rows_per_user": current + future,
            "mandatory_database_rows_per_user": 1,
        }
    )
    with pytest.raises(ProductError) as pressure:
        service.create_api_key(user_id=user_id, name="must wait")
    assert pressure.value.code == "database_quota_reached"
    for recovery in range(service.settings.max_recovery_attempts):
        claim = service.claim_next_job(f"row-pressure-recovery-{recovery}")
        assert claim is not None
        with service.database.transaction(immediate=True) as conn:
            conn.execute(
                "UPDATE jobs SET lease_expires_at=? WHERE id=?",
                (timestamp(utcnow() - timedelta(seconds=1)), claim.job_id),
            )
        assert service.recover_interrupted_jobs() == 1
    _install_successful_extractor(service)
    assert service.process_one("row-pressure-worker")
    completed = service.get_job(user_id=user_id, job_id=str(job["id"]))
    assert completed["status"] == "review"
    with service.database.connect() as conn:
        assert (
            service._database_row_count(conn, user_id)
            <= service.settings.max_database_rows_per_user
        )
        assert conn.execute("SELECT COUNT(*) FROM provider_attempts").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM result_versions").fetchone()[0] == 1


def test_free_reprocess_admission_is_atomic_and_preserves_existing_terminal_capacity(
    tmp_path: Path,
) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=1)
    user_id = customer_id(service, "free-reprocess-pressure@example.com")
    fixture_upload = service.prepare_demo_upload(user_id)
    free_job = service.create_job(
        user_id=user_id,
        upload_id=str(fixture_upload["id"]),
        page_index=0,
        crop=None,
    )
    _install_successful_extractor(service)
    assert service.process_one("free-fixture-worker")
    queued = _paid_job(service, user_id, color="navy")

    with service.database.connect() as conn:
        current = service._database_row_count(conn, user_id)
        future = service._future_terminal_rows(conn, user_id)
        before_job = dict(
            conn.execute("SELECT * FROM jobs WHERE id=?", (free_job["id"],)).fetchone()
        )
        before_counts = {
            table: conn.execute(
                f'SELECT COUNT(*) FROM "{table}" WHERE user_id=?',  # noqa: S608
                (user_id,),
            ).fetchone()[0]
            for table in (
                "audit_events",
                "credit_ledger",
                "provider_attempts",
                "result_versions",
                "storage_reservations",
            )
        }
    assert future > 0
    service.settings = Settings(
        **{
            **service.settings.__dict__,
            "max_database_rows_per_user": current + future + 1,
            "mandatory_database_rows_per_user": 1,
        }
    )

    with pytest.raises(ProductError) as rejected:
        service.reprocess(user_id=user_id, job_id=str(free_job["id"]))
    assert rejected.value.code == "database_quota_reached"
    with service.database.connect() as conn:
        after_job = dict(
            conn.execute("SELECT * FROM jobs WHERE id=?", (free_job["id"],)).fetchone()
        )
        assert after_job == before_job
        after_counts = {
            table: conn.execute(
                f'SELECT COUNT(*) FROM "{table}" WHERE user_id=?',  # noqa: S608
                (user_id,),
            ).fetchone()[0]
            for table in before_counts
        }
    assert after_counts == before_counts
    assert service.account(user_id)["credits"] == 0

    # The previously admitted queued job still has every row it needs to reach a
    # durable terminal result under the same pressure ceiling.
    assert service.process_one("existing-capacity-worker")
    assert service.get_job(user_id=user_id, job_id=str(queued["id"]))["status"] == "review"


def test_invalid_provider_output_never_logs_success_and_uses_specific_error(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=1)
    user_id = customer_id(service, "invalid-provider@example.com")
    job = _paid_job(service, user_id, color="blue")
    service.extractor = SimpleNamespace(
        extract=lambda _: SimpleNamespace(
            chart=SimpleNamespace(chart_type="line", series=[]),
            raw="invalid",
            extractor="invalid-provider",
            model_version="invalid-release",
        )
    )
    with caplog.at_level("INFO"):
        assert service.process_one("invalid-output-worker")
    failed = service.get_job(user_id=user_id, job_id=str(job["id"]))
    assert failed["status"] == "failed"
    assert failed["error"]["code"] == "model_output_invalid"
    assert not any(record.message == "provider_call_succeeded" for record in caplog.records)
    with service.database.connect() as conn:
        attempt = conn.execute("SELECT outcome FROM provider_attempts").fetchone()
    assert attempt["outcome"] == "failed"


def test_logout_revokes_all_sessions_and_rotates_principal_generation(tmp_path: Path) -> None:
    service = service_for(tmp_path, seed_demo_account=False)
    first = service.register("sessions@example.com", "long session password")
    user = service.session_user(first["session"])
    assert user is not None
    user_id = str(user["id"])
    second = service.create_session(user_id)
    before = service.account(user_id)["principal_marker"]
    service.logout(first["session"])
    assert service.session_user(first["session"]) is None
    assert service.session_user(second["session"]) is None
    after = service.account(user_id)["principal_marker"]
    assert after != before
    replacement = service.create_session(user_id)
    assert service.session_user(replacement["session"])["id"] == user_id


def test_backup_and_restore_fsync_files_manifests_and_publication_directories(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = service_for(tmp_path, seed_demo_account=False, initial_credits=1)
    user_id = customer_id(service, "fsync-backup@example.com")
    _paid_job(service, user_id, color="orange")
    directories: list[Path] = []
    files: list[Path] = []
    real_directory = backup_module._fsync_directory
    real_file = backup_module._fsync_file

    def tracked_directory(path: Path) -> None:
        directories.append(path)
        real_directory(path)

    def tracked_file(path: Path) -> None:
        files.append(path)
        real_file(path)

    monkeypatch.setattr(backup_module, "_fsync_directory", tracked_directory)
    monkeypatch.setattr(backup_module, "_fsync_file", tracked_file)
    reservation = service._reserve_storage_bytes(
        user_id=user_id,
        byte_count=1,
        kind="staging",
        storage_path=service.storage.root / "staging" / "in-flight.tmp",
    )
    blocked_destination = tmp_path / "blocked-in-flight-backup"
    with pytest.raises(BackupError, match="publication is in flight"):
        create_backup(service.settings, blocked_destination)
    assert not blocked_destination.exists()
    service._release_storage_reservation(user_id=user_id, reservation=reservation)
    destination = tmp_path / "fsync-backup"
    create_backup(service.settings, destination)
    assert directories.count(destination.parent.resolve()) >= 2
    assert any(path.name == "manifest.json" for path in files)
    directories.clear()
    files.clear()
    target = tmp_path / "fsync-restore"
    restore_backup(destination, target)
    assert directories.count(target.parent.resolve()) >= 2
    assert any(path.name == "unrender.sqlite3" for path in files)


@pytest.mark.parametrize("script", ["browser_auth_epoch.mjs", "browser_two_tab.mjs"])
def test_browser_privacy_editor_and_two_tab_regressions(script: str) -> None:
    result = subprocess.run(
        ["node", str(Path(__file__).with_name(script))],
        cwd=STATIC_DIR.parents[2],
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr or result.stdout


def test_modal_deadline_drains_worker_without_refund_or_redispatch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    started = threading.Event()
    cancelled = threading.Event()
    calls: list[bytes] = []

    async def remote(image: bytes, *_: object) -> None:
        calls.append(image)
        started.set()
        try:
            await asyncio.Event().wait()
        finally:
            cancelled.set()

    monkeypatch.setitem(
        sys.modules,
        "modal",
        SimpleNamespace(
            Function=SimpleNamespace(
                from_name=lambda *_: SimpleNamespace(remote=SimpleNamespace(aio=remote))
            )
        ),
    )
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        initial_credits=1,
        extractor_backend="modal",
        provider_timeout_seconds=1,
    )
    service.extractor = ModalExtractor(service.settings)
    user_id = customer_id(service)
    job = _paid_job(service, user_id, color="purple")
    worker = JobWorker(service, poll_seconds=0.01)
    worker.start()
    try:
        assert started.wait(timeout=5)
        assert worker.stop(timeout=5)
    finally:
        worker.stop(timeout=5)
    assert cancelled.is_set()
    result = service.get_job(user_id=user_id, job_id=str(job["id"]))
    assert result["status"] == "failed"
    assert result["error"]["code"] == "provider_timeout"
    assert service.account(user_id)["credits"] == 0
    assert service.recover_interrupted_jobs() == 0
    assert not service.process_one("after-restart")
    assert len(calls) == 1
    with service.database.connect() as conn:
        assert (
            conn.execute(
                "SELECT provider_dispatched FROM jobs WHERE id=?", (job["id"],)
            ).fetchone()[0]
            == 1
        )


def test_modal_canary_bounds_lazy_resolution_without_inference(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def hydrate() -> None:
        await asyncio.Event().wait()

    def remote(*_: object) -> None:
        pytest.fail("Canary must not invoke inference")

    monkeypatch.setitem(
        sys.modules,
        "modal",
        SimpleNamespace(
            Function=SimpleNamespace(
                from_name=lambda *_: SimpleNamespace(
                    hydrate=SimpleNamespace(aio=hydrate), remote=SimpleNamespace(aio=remote)
                )
            )
        ),
    )
    extractor = ModalExtractor(settings_for(tmp_path, provider_timeout_seconds=1))
    with pytest.raises(ExtractionError) as error:
        extractor.canary_contract()
    assert error.value.code == "provider_contract_unavailable"


@pytest.mark.parametrize("deadline", [0, 241])
def test_provider_deadline_rejects_unbounded_configuration(tmp_path: Path, deadline: int) -> None:
    with pytest.raises(ValueError, match="PROVIDER_TIMEOUT_SECONDS"):
        settings_for(tmp_path, provider_timeout_seconds=deadline).validate()


def test_concurrent_logins_can_share_a_legacy_password_upgrade(tmp_path, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor

    service = service_for(tmp_path, seed_demo_account=False)
    password = "legacy password is long enough"
    user_id = service.provision_user("legacy@example.com", password)
    salt = b"legacy-salt-1234"
    derived = hashlib.scrypt(password.encode(), salt=salt, n=2**14, r=8, p=1)
    legacy = "scrypt$16384$8$1$" + base64.urlsafe_b64encode(salt).decode()
    legacy += "$" + base64.urlsafe_b64encode(derived).decode()
    with service.database.transaction(immediate=True) as conn:
        conn.execute("UPDATE users SET password_hash=? WHERE id=?", (legacy, user_id))
    barrier = threading.Barrier(2)
    original = service_module.verify_password

    def overlapping_verification(value, encoded):
        result = original(value, encoded)
        if encoded == legacy:
            barrier.wait(timeout=5)
        return result

    monkeypatch.setattr(service_module, "verify_password", overlapping_verification)
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [
            pool.submit(service.authenticate, "legacy@example.com", password) for _ in range(2)
        ]
        sessions = [future.result(timeout=10) for future in futures]
    assert len({session["session"] for session in sessions}) == 2
    assert all(service.session_user(session["session"]) for session in sessions)


def test_worker_recovers_lease_that_expires_after_startup_without_redispatch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = service_for(
        tmp_path,
        seed_demo_account=False,
        initial_credits=1,
        worker_lease_seconds=10,
        worker_heartbeat_seconds=1,
    )
    user_id = customer_id(service)
    job = _paid_job(service, user_id, color="purple")
    claim = service.claim_next_job("terminated-worker")
    assert claim is not None and service._begin_provider_dispatch(claim)
    service.initialize()
    assert service.get_job(user_id=user_id, job_id=str(job["id"]))["status"] == "running"

    polled = threading.Event()
    original_process = service.process_one
    provider_calls: list[bytes] = []

    def observe_poll(*args: Any, **kwargs: Any) -> bool:
        result = original_process(*args, **kwargs)
        polled.set()
        return result

    def unexpected_provider(image: bytes) -> ExtractionOutput:
        provider_calls.append(image)
        raise AssertionError("Dispatched work must not be invoked again")

    monkeypatch.setattr(service, "process_one", observe_poll)
    service.extractor = SimpleNamespace(extract=unexpected_provider)
    worker = JobWorker(service, poll_seconds=0.01)
    worker.start()
    try:
        assert polled.wait(timeout=5)
        with service.database.transaction(immediate=True) as conn:
            conn.execute(
                "UPDATE jobs SET lease_expires_at=? WHERE id=?",
                (timestamp(utcnow() - timedelta(seconds=1)), job["id"]),
            )
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            current = service.get_job(user_id=user_id, job_id=str(job["id"]))
            if current["status"] == "failed":
                break
            time.sleep(0.02)
        assert current["status"] == "failed"
        assert current["error"]["code"] == "worker_lease_expired_after_dispatch"
        assert service.account(user_id)["credits"] == 0
        assert provider_calls == []
        with service.database.connect() as conn:
            assert (
                conn.execute(
                    "SELECT COUNT(*) FROM provider_attempts WHERE job_id=?", (job["id"],)
                ).fetchone()[0]
                == 1
            )
    finally:
        assert worker.stop(timeout=5)


def test_server_entrypoint_emits_private_structured_lifecycle_logs() -> None:
    script = """
import logging
import logging.config
import uvicorn
from unrender.product.cli import main

def run(*args, **kwargs):
    logging.config.dictConfig(kwargs['log_config'])
    logger = logging.getLogger('unrender.product')
    logger.info('provider_call_succeeded', extra={
        'job_id': 'test-job', 'provider': 'modal', 'duration_ms': 123,
        'source_filename': 'private-chart.png', 'password': 'private-password',
    })
    try:
        raise ValueError('private exception detail')
    except ValueError:
        logger.exception('job_processing_failed', extra={'job_id': 'test-job'})

uvicorn.run = run
main()
"""
    result = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    records = [json.loads(line) for line in result.stdout.splitlines()]
    assert len(records) == 2
    assert records[0]["event"] == "provider_call_succeeded"
    assert records[0]["provider"] == "modal"
    assert records[0]["duration_ms"] == 123
    assert records[1]["event"] == "job_processing_failed"
    assert records[1]["exception_type"] == "ValueError"
    assert all(record["job_id"] == "test-job" for record in records)
    assert "private" not in result.stdout + result.stderr
    assert "Traceback" not in result.stdout + result.stderr
